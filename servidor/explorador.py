"""As abas Arquivos e Visualizador do modo Code: a árvore da pasta do trabalho e o que cada arquivo tem dentro.

Toda rota recebe `raiz` (a pasta da conversa, caminho completo) e `caminho` (relativo a ela, com / ou \\). Nada sai
da raiz: caminho absoluto, `..`, `:` (unidade ou fluxo alternativo do NTFS) e link ou junção que aponte para fora são
recusados. Os arquivos protegidos (cfg["protegidos"], os mesmos das ferramentas do Code) só mostram o conteúdo com
`forcar=1`, que a tela manda depois do clique em "Mostrar mesmo assim".

Rotas (/api/arquivos/...): lista, busca, conteudo, bruto, planilha, docx, info, mtimes, mudancas (GET ou POST de
leitura) e criar, renomear, lixeira, abrir, revelar (POST).
"""
from __future__ import annotations

import codecs
import collections
import csv
import ctypes
import datetime as dt
import difflib
import fnmatch
import io
import mimetypes
import os
import re
import sqlite3
import stat
import subprocess
import sys
import time
import urllib.parse
from pathlib import Path

import banco
import config
from i18n import t
from rotas import Erro, rota

LIMITE_TEXTO = 2 * 1024 * 1024
LIMITE_ITENS = 2000
LIMITE_BUSCA = 200
TEMPO_BUSCA = 4.0
LINHAS_PLANILHA = 200
COLUNAS_PLANILHA = 50
ABAS_PLANILHA = 40

# escondidas por padrão na árvore (ocultos=1 mostra)
OCULTAS = {".git", "node_modules", "__pycache__", ".venv", "dist"}
# a busca nunca entra nelas: pesadas e sem nada que a pessoa procure pelo nome
PESADAS = OCULTAS | {"venv", ".mypy_cache", ".pytest_cache", ".next", ".cache", ".idea", ".vs", ".tox", ".gradle",
                     "build", "target", "Library", "obj"}

# abrir com o programa padrão executa estes (ou muda o sistema, ou roda macro): a tela pede confirmação antes.
# A base é a lista de anexos que o Outlook bloqueia, mais os formatos do Office com macro e os scripts que um
# programa instalado costuma abrir com dois cliques.
EXECUTAVEIS = {
    # programas, instaladores e atalhos
    ".exe", ".com", ".scr", ".pif", ".cpl", ".msi", ".msp", ".mst", ".msu", ".msix", ".msixbundle", ".appx",
    ".appxbundle", ".appinstaller", ".application", ".appref-ms", ".xbap", ".jar", ".jnlp", ".gadget", ".vsix",
    ".vsto", ".lnk", ".url", ".website", ".scf", ".library-ms", ".searchconnector-ms", ".settingcontent-ms",
    ".appcontent-ms", ".msc", ".diagcab", ".diagcfg", ".diagpkg", ".wsb", ".xnk", ".prf", ".ops",
    # muda o sistema ao abrir
    ".reg", ".inf", ".ins", ".isp", ".theme", ".themepack", ".deskthemepack",
    # scripts do Windows
    ".bat", ".cmd", ".vbs", ".vbe", ".vb", ".vbp", ".bas", ".js", ".jse", ".ws", ".wsf", ".wsh", ".wsc", ".sct",
    ".hta", ".shb", ".shs", ".chm", ".hlp",
    # PowerShell (e o antigo Monad)
    ".ps1", ".psm1", ".psd1", ".ps1xml", ".ps2", ".ps2xml", ".psc1", ".psc2", ".pssc", ".psrc", ".cdxml", ".msh",
    ".msh1", ".msh2", ".mshxml", ".msh1xml", ".msh2xml",
    # outras linguagens que rodam com dois cliques quando estão instaladas
    ".py", ".pyw", ".pyc", ".pyo", ".pyz", ".pyzw", ".sh", ".pl", ".rb", ".rbw", ".ahk", ".au3",
    # Office com macro (e os suplementos)
    ".docm", ".dotm", ".xlsm", ".xltm", ".xlsb", ".xla", ".xlam", ".xll", ".xlm", ".pptm", ".potm", ".ppsm",
    ".ppam", ".ppa", ".sldm", ".vsmacros",
    # Access (roda a macro de abertura)
    ".mdb", ".mde", ".mda", ".accdb", ".accde", ".accdr", ".accda", ".ade", ".adp",
}

_IDIOMAS = {
    ".py": "python", ".pyw": "python", ".pyi": "python", ".js": "javascript", ".mjs": "javascript",
    ".cjs": "javascript", ".jsx": "javascript", ".ts": "typescript", ".tsx": "typescript", ".mts": "typescript",
    ".json": "json", ".jsonc": "json", ".json5": "json", ".webmanifest": "json", ".ipynb": "json",
    ".html": "xml", ".htm": "xml", ".xhtml": "xml", ".xml": "xml", ".svg": "xml", ".xaml": "xml", ".csproj": "xml",
    ".vbproj": "xml", ".props": "xml", ".targets": "xml", ".plist": "xml", ".config": "xml", ".resx": "xml",
    ".css": "css", ".scss": "scss", ".sass": "scss", ".less": "less", ".md": "markdown", ".markdown": "markdown",
    ".mdx": "markdown", ".yml": "yaml", ".yaml": "yaml", ".toml": "ini", ".ini": "ini", ".cfg": "ini", ".conf": "ini",
    ".properties": "ini", ".env": "bash", ".sh": "bash", ".bash": "bash", ".zsh": "bash", ".ps1": "powershell",
    ".psm1": "powershell", ".psd1": "powershell", ".bat": "dos", ".cmd": "dos", ".c": "c", ".h": "c", ".cpp": "cpp",
    ".cc": "cpp", ".cxx": "cpp", ".hpp": "cpp", ".hh": "cpp", ".ino": "cpp", ".cs": "csharp", ".java": "java",
    ".kt": "kotlin", ".kts": "kotlin", ".go": "go", ".rs": "rust", ".rb": "ruby", ".php": "php", ".lua": "lua",
    ".sql": "sql", ".swift": "swift", ".r": "r", ".pl": "perl", ".pm": "perl", ".vb": "vbnet", ".vbs": "vbnet",
    ".m": "objectivec", ".mm": "objectivec", ".diff": "diff", ".patch": "diff", ".graphql": "graphql",
    ".gql": "graphql", ".mk": "makefile", ".gml": "javascript", ".wat": "wasm", ".csv": "plaintext",
    ".tsv": "plaintext", ".txt": "plaintext", ".log": "plaintext",
}
_IDIOMAS_NOME = {"makefile": "makefile", "gnumakefile": "makefile", "dockerfile": "bash", ".gitignore": "bash",
                 ".gitattributes": "bash", ".editorconfig": "ini", ".npmrc": "ini", ".env": "bash",
                 "requirements.txt": "plaintext", "cmakelists.txt": "makefile"}

_MIME_EXTRA = {".webp": "image/webp", ".avif": "image/avif", ".ico": "image/x-icon", ".svg": "image/svg+xml",
               ".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".aac": "audio/aac", ".flac": "audio/flac",
               ".ogg": "audio/ogg", ".oga": "audio/ogg", ".opus": "audio/ogg", ".wav": "audio/wav",
               ".weba": "audio/webm", ".mp4": "video/mp4", ".m4v": "video/mp4", ".webm": "video/webm",
               ".mov": "video/quicktime", ".ogv": "video/ogg", ".mkv": "video/x-matroska", ".pdf": "application/pdf",
               ".md": "text/markdown", ".csv": "text/csv"}
# o que o bruto nunca entrega com o tipo de verdade: o navegador poderia montar uma página com ele
_VIRA_TEXTO = ("text/html", "application/xhtml+xml", "text/xml", "application/xml", "text/javascript",
               "application/javascript", "application/x-javascript", "text/css", "application/json")
# sem scripts nem nada de fora; o leitor de PDF do Electron 44 e do Edge abre normalmente com ele (conferido)
_CSP_BRUTO = "sandbox; default-src 'none'; img-src 'self' data:; media-src 'self'; style-src 'unsafe-inline'"

_NOME_RESERVADO = re.compile(r"^(con|prn|aux|nul|com[0-9¹²³]|lpt[0-9¹²³])(\..*)?$", re.I)
_PROIBIDOS = set('<>:"/\\|?*')


# ======================= caminhos =======================

def _no_windows() -> bool:
    return sys.platform == "win32"


def _raiz(valor) -> Path:
    valor = str(valor or "").strip()
    if not valor:
        raise Erro(400, t("Falta dizer qual é a pasta do trabalho."))
    p = Path(valor)
    if not p.is_absolute():
        raise Erro(400, t("A pasta do trabalho precisa ser um caminho completo (como C:\\Projetos\\site)."))
    try:
        p = p.resolve(strict=True)
    except (OSError, RuntimeError):
        raise Erro(404, t("A pasta do trabalho não existe mais: {caminho}", caminho=valor))
    if not p.is_dir():
        raise Erro(404, t("A pasta do trabalho não existe mais: {caminho}", caminho=valor))
    return p


def _partes(caminho) -> list[str]:
    c = str(caminho or "").replace("\\", "/").strip()
    if c in ("", "."):
        return []
    if c.startswith("/") or re.match(r"^[A-Za-z]:", c):
        raise Erro(400, t("Use um caminho dentro da pasta do trabalho, não um caminho completo."))
    partes = [p for p in c.split("/") if p not in ("", ".")]
    for p in partes:
        if p == "..":
            raise Erro(403, t("Esse caminho sai da pasta do trabalho (..), então não abre aqui."))
        if ":" in p or "\x00" in p:
            raise Erro(400, t("Caminho inválido: {caminho}", caminho=caminho))
    return partes


def _dentro(raiz: Path, p: Path) -> bool:
    try:
        r = p.resolve()
    except (OSError, RuntimeError):
        return False
    return r == raiz or r.is_relative_to(raiz)


def _e_link(p: Path) -> bool:
    try:
        return p.is_symlink() or (hasattr(os.path, "isjunction") and os.path.isjunction(p))
    except OSError:
        return False


def _local(raiz: Path, caminho, *, existir=True, seguir=True) -> Path:
    """O caminho dentro da raiz (sem resolver links no último pedaço). `seguir`: o destino real também precisa estar
    dentro (ler, listar); sem ele, só a pasta onde o item mora precisa (renomear e mandar para a Lixeira um link)."""
    partes = _partes(caminho)
    p = raiz.joinpath(*partes)
    if partes and not _dentro(raiz, p.parent):
        raise Erro(403, t("Esse caminho passa por um atalho que aponta para fora da pasta do trabalho."))
    if existir and not (p.exists() or _e_link(p)):
        raise Erro(404, t("Não achei {caminho}. Ele pode ter sido apagado ou movido.", caminho=_rel(raiz, p)))
    if seguir and (p.exists() or not existir) and not _dentro(raiz, p):
        raise Erro(403, t("Este atalho aponta para fora da pasta do trabalho, então não abre aqui."))
    return p


def _rel(raiz: Path, p: Path) -> str:
    try:
        return p.relative_to(raiz).as_posix()
    except ValueError:
        return p.name


def _cfg_protegidos() -> list[str]:
    return list(config.carregar().get("protegidos") or [])


def _protegido(raiz: Path, p: Path, protegidos: list[str] | None = None) -> str | None:
    from ferramentas import code as fcode  # aqui dentro: o code.py puxa o resto das ferramentas
    ctx = fcode.Contexto(raiz=raiz, protegidos=_cfg_protegidos() if protegidos is None else protegidos,
                         execucao_id="", conversa_id="")
    return ctx.protegido(p)


def _exigir_liberado(h, raiz: Path, p: Path, q: dict) -> bool:
    """Arquivo protegido só passa com forcar=1. Responde e devolve False quando não pode."""
    padrao = _protegido(raiz, p)
    if padrao and str(q.get("forcar") or "") not in ("1", "true"):
        h._json({"erro": t("Este arquivo está na lista de protegidos ({padrao}): ele pode ter senhas ou chaves.",
                           padrao=padrao), "codigo": "protegido", "padrao": padrao}, 403)
        return False
    return True


def _extensao(nome: str) -> str:
    """A extensão como o Windows vê: ponto e espaço no fim do nome não contam (o Windows tira e abre "roda.bat." como
    "roda.bat"), e um nome que começa com ponto (".bat") também tem extensão."""
    nome = str(nome).rstrip(". ")
    i = nome.rfind(".")
    return nome[i:].lower() if i >= 0 else ""


def _executavel(p: Path) -> bool:
    """Abrir com o programa padrão executa isto? Confere o nome pedido e o nome de verdade do arquivo (o nome curto
    PROGRA~1.APP de "programa.appref-ms", um nome com ponto no fim e um link apontam para outro nome)."""
    nomes = {p.name}
    try:
        nomes.add(p.resolve().name)
    except (OSError, RuntimeError):
        pass
    return any(_extensao(n) in EXECUTAVEIS for n in nomes)


def _stat_entrada(entrada: os.DirEntry):
    """O stat de um item da pasta (o do destino, num link; o do próprio link quando o destino não abre)."""
    try:
        return entrada.stat()
    except OSError:
        try:
            return entrada.stat(follow_symlinks=False)
        except OSError:
            return None


def _oculto_windows(st) -> bool:
    atrib = getattr(st, "st_file_attributes", 0)
    return bool(atrib & stat.FILE_ATTRIBUTE_HIDDEN and atrib & stat.FILE_ATTRIBUTE_SYSTEM) if atrib else False


def _escondido(nome: str, st) -> bool:
    """O que a árvore esconde sem "mostrar as pastas de sistema": .git, node_modules... e o que o Windows marca como
    oculto e de sistema ao mesmo tempo (a busca segue a mesma regra)."""
    return nome in OCULTAS or (st is not None and _oculto_windows(st))


def _chave_natural(nome: str):
    return [int(p) if p.isdigit() else p.casefold() for p in re.split(r"(\d+)", nome)]


def _validar_nome(nome) -> str:
    nome = str(nome or "").strip()
    if not nome or nome in (".", ".."):
        raise Erro(400, t("Escreva um nome."))
    if any(c in _PROIBIDOS for c in nome) or any(ord(c) < 32 for c in nome):
        raise Erro(400, t("O nome não pode ter estes caracteres: {lista}", lista='< > : " / \\ | ? *'))
    if nome.endswith((".", " ")):
        raise Erro(400, t("O nome não pode terminar com ponto ou espaço."))
    if _NOME_RESERVADO.match(nome):
        raise Erro(400, t("O Windows reserva o nome {nome}. Escolha outro.", nome=nome))
    if len(nome) > 255:
        raise Erro(400, t("Nome comprido demais."))
    return nome


# ======================= texto =======================

_CONTROLES = bytes(c for c in range(32) if c not in (8, 9, 10, 12, 13, 27))


def _parece_utf16(amostra: bytes) -> str | None:
    n = len(amostra) // 2 * 2
    if n < 8:
        return None
    pares, impares = amostra[0:n:2], amostra[1:n:2]
    zeros_par, zeros_impar = pares.count(0) / len(pares), impares.count(0) / len(impares)
    if zeros_impar > 0.3 and zeros_par < 0.02:
        return "utf-16-le"
    if zeros_par > 0.3 and zeros_impar < 0.02:
        return "utf-16-be"
    return None


def binario(amostra: bytes) -> bool:
    if not amostra:
        return False
    if b"\x00" in amostra:
        return True
    controles = sum(amostra.count(c) for c in _CONTROLES)
    return controles / len(amostra) > 0.1


def decodificar(dados: bytes, cortado: bool = False) -> tuple[str, str] | None:
    """(texto, codificação), ou None quando o arquivo parece binário. Com `cortado`, o fim pode ter um caractere
    pela metade, que fica de fora."""
    for bom, cod, nome in ((codecs.BOM_UTF16_LE, "utf-16-le", "utf-16"), (codecs.BOM_UTF16_BE, "utf-16-be", "utf-16")):
        if dados.startswith(bom) and not dados.startswith(codecs.BOM_UTF32_LE):
            corpo = dados[2:]
            if len(corpo) % 2:
                corpo = corpo[:-1]
            return corpo.decode(cod, errors="replace"), nome
    sem_bom = _parece_utf16(dados[:4096])
    if sem_bom:
        corpo = dados[: len(dados) // 2 * 2]
        return corpo.decode(sem_bom, errors="replace"), "utf-16"
    com_bom = dados.startswith(codecs.BOM_UTF8)
    corpo = dados[3:] if com_bom else dados
    if binario(corpo[:8192]):
        return None
    try:
        texto = codecs.getincrementaldecoder("utf-8")("strict").decode(corpo, final=not cortado)
        return texto, "utf-8-sig" if com_bom else "utf-8"
    except UnicodeDecodeError:
        return corpo.decode("cp1252", errors="replace"), "cp1252"


def linguagem(p: Path) -> str:
    nome = p.name.lower()
    if nome in _IDIOMAS_NOME:
        return _IDIOMAS_NOME[nome]
    if nome.startswith(".env"):
        return "bash"
    return _IDIOMAS.get(p.suffix.lower(), "plaintext")


def _mime(p: Path) -> str:
    ext = p.suffix.lower()
    return _MIME_EXTRA.get(ext) or mimetypes.guess_type(p.name)[0] or "application/octet-stream"


def _ler_inicio(p: Path, limite: int) -> tuple[bytes, int]:
    tamanho = p.stat().st_size
    with open(p, "rb") as f:
        return f.read(limite), tamanho


# ======================= rotas de leitura =======================

def _item(raiz: Path, entrada: os.DirEntry, protegidos: list[str]) -> dict | None:
    p = Path(entrada.path)
    link = entrada.is_symlink() or (hasattr(entrada, "is_junction") and entrada.is_junction())
    st = _stat_entrada(entrada)
    if st is None:
        return None
    try:
        pasta = entrada.is_dir()
    except OSError:
        pasta = False
    d = {"nome": entrada.name, "caminho": _rel(raiz, p), "tipo": "pasta" if pasta else "arquivo",
         "modificado": st.st_mtime}
    if not pasta:
        d["tamanho"] = st.st_size
    if link:
        d["link"] = True
        if not _dentro(raiz, p):
            d["fora"] = True
    padrao = _protegido(raiz, p, protegidos)
    if padrao:
        d["protegido"] = padrao
    if _escondido(entrada.name, st):
        d["oculto"] = True
    return d


@rota("GET", "/api/arquivos/lista")
def r_lista(h):
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    pasta = _local(raiz, q.get("caminho"))
    if not pasta.is_dir():
        raise Erro(400, t("{caminho} não é uma pasta.", caminho=_rel(raiz, pasta)))
    mostrar_ocultos = str(q.get("ocultos") or "") in ("1", "true")
    protegidos = _cfg_protegidos()
    itens, escondidos = [], 0
    try:
        with os.scandir(pasta) as it:
            entradas = list(it)
    except PermissionError:
        raise Erro(403, t("O Windows não deixou abrir esta pasta (sem permissão)."))
    except OSError as e:
        raise Erro(400, t("Não deu para abrir esta pasta: {erro}", erro=e.strerror or e))
    for entrada in entradas:
        d = _item(raiz, entrada, protegidos)
        if d is None:
            continue
        if d.get("oculto") and not mostrar_ocultos:
            escondidos += 1
            continue
        itens.append(d)
    itens.sort(key=lambda d: (d["tipo"] != "pasta", _chave_natural(d["nome"])))
    total = len(itens)
    try:
        modificado = pasta.stat().st_mtime
    except OSError:
        modificado = 0
    h._json({"caminho": _rel(raiz, pasta) if pasta != raiz else "", "itens": itens[:LIMITE_ITENS], "total": total,
             "cortado": total > LIMITE_ITENS, "escondidos": escondidos, "modificado": modificado})


@rota("GET", "/api/arquivos/busca")
def r_busca(h):
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    termo = str(q.get("termo") or "").strip()
    if not termo:
        raise Erro(400, t("Escreva o que procurar."))
    mostrar_ocultos = str(q.get("ocultos") or "") in ("1", "true")
    padrao = termo.casefold()
    glob = any(c in termo for c in "*?[")
    resultados, inicio, esgotou = [], time.monotonic(), False
    fila = collections.deque([raiz])
    while fila and len(resultados) < LIMITE_BUSCA:
        pasta = fila.popleft()
        try:
            with os.scandir(pasta) as it:
                entradas = sorted(it, key=lambda e: _chave_natural(e.name))
        except OSError:
            continue
        for e in entradas:
            if time.monotonic() - inicio > TEMPO_BUSCA:
                esgotou = True
                break
            nome = e.name
            try:
                eh_pasta = e.is_dir(follow_symlinks=False)
            except OSError:
                continue
            link = e.is_symlink() or (hasattr(e, "is_junction") and e.is_junction())
            # o que a árvore esconde (de sistema), a busca também esconde, e não entra nas pastas assim
            visivel = mostrar_ocultos or not _escondido(nome, _stat_entrada(e))
            nome_cf = nome.casefold()
            achou = fnmatch.fnmatchcase(nome_cf, padrao) if glob else padrao in nome_cf
            if achou and visivel:
                resultados.append({"nome": nome, "caminho": _rel(raiz, Path(e.path)),
                                   "tipo": "pasta" if (eh_pasta or (link and os.path.isdir(e.path))) else "arquivo"})
                if len(resultados) >= LIMITE_BUSCA:
                    break
            if eh_pasta and not link and nome not in PESADAS and visivel:
                fila.append(Path(e.path))
        if esgotou:
            break
    h._json({"resultados": resultados, "cortado": len(resultados) >= LIMITE_BUSCA, "tempo_esgotado": esgotou})


@rota("GET", "/api/arquivos/conteudo")
def r_conteudo(h):
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    p = _local(raiz, q.get("caminho"))
    if p.is_dir():
        raise Erro(400, t("{caminho} é uma pasta.", caminho=_rel(raiz, p)))
    if not _exigir_liberado(h, raiz, p, q):
        return
    try:
        dados, tamanho = _ler_inicio(p, LIMITE_TEXTO)
        modificado = p.stat().st_mtime
    except PermissionError:
        raise Erro(403, t("O Windows não deixou ler este arquivo (sem permissão ou aberto em outro programa)."))
    except OSError as e:
        raise Erro(400, t("Não deu para ler o arquivo: {erro}", erro=e.strerror or e))
    cortado = tamanho > len(dados)
    base = {"caminho": _rel(raiz, p), "nome": p.name, "tamanho": tamanho, "modificado": modificado,
            "mime": _mime(p), "cortado": cortado}
    r = decodificar(dados, cortado)
    if r is None:
        return h._json({**base, "binario": True})
    texto, cod = r
    h._json({**base, "binario": False, "texto": texto, "codificacao": cod, "linguagem": linguagem(p),
             "crlf": "\r\n" in texto[:20000]})


def _tipo_bruto(p: Path) -> str:
    tipo = _mime(p)
    if tipo in _VIRA_TEXTO or tipo.startswith("text/") or tipo.endswith(("+xml", "+json")) and tipo != "image/svg+xml":
        return "text/plain"
    if tipo.startswith(("image/", "audio/", "video/")) or tipo == "application/pdf":
        return tipo
    return "application/octet-stream"


@rota("GET", "/api/arquivos/bruto")
@rota("GET", "/api/arquivos/bruto/<nome>")
def r_bruto(h, nome=None):
    """O arquivo como ele é (imagem, PDF, áudio, vídeo), com Range. O /<nome> no fim do endereço é só para o leitor de
    PDF mostrar o nome do arquivo no título; o que vale é o `caminho` da query."""
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    p = _local(raiz, q.get("caminho"))
    if not p.is_file():
        raise Erro(400, t("{caminho} não é um arquivo.", caminho=_rel(raiz, p)))
    if not _exigir_liberado(h, raiz, p, q):
        return
    tipo = _tipo_bruto(p)
    try:
        total = p.stat().st_size
        f = open(p, "rb")
    except OSError as e:
        raise Erro(403 if isinstance(e, PermissionError) else 400,
                   t("Não deu para ler o arquivo: {erro}", erro=e.strerror or e))
    with f:
        cab = {"Accept-Ranges": "bytes", "Cache-Control": "no-cache", "X-Content-Type-Options": "nosniff",
               "Referrer-Policy": "no-referrer", "Content-Security-Policy": _CSP_BRUTO,
               "Content-Disposition": "inline; filename*=UTF-8''" + urllib.parse.quote(p.name)}
        ini, fim, status = 0, total - 1, 200
        m = re.match(r"^bytes=(\d*)-(\d*)$", (h.headers.get("Range") or "").strip())
        if m and (m.group(1) or m.group(2)):
            if m.group(1):
                ini, fim = int(m.group(1)), min(int(m.group(2) or total - 1), total - 1)
            else:  # bytes=-N: os últimos N
                ini = max(0, total - int(m.group(2)))
            if ini > fim or ini >= total:
                return h._bytes(b"", tipo, status=416, cab={**cab, "Content-Range": f"bytes */{total}"})
            status = 206
            cab["Content-Range"] = f"bytes {ini}-{fim}/{total}"
        tamanho = max(0, fim - ini + 1)
        h.send_response(status)
        h.send_header("Content-Type", tipo)
        h.send_header("Content-Length", str(tamanho))
        for k, v in cab.items():
            h.send_header(k, v)
        h.end_headers()
        if h.command == "HEAD":
            return
        f.seek(ini)
        falta = tamanho
        while falta > 0:
            bloco = f.read(min(256 * 1024, falta))
            if not bloco:
                break
            h.wfile.write(bloco)
            falta -= len(bloco)


def _celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return t("VERDADEIRO") if v else t("FALSO")
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:.12g}"
    return str(v)


def _planilha_xlsx(p: Path) -> list[dict]:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        wb_v = openpyxl.load_workbook(p, read_only=True, data_only=True)
        wb_f = openpyxl.load_workbook(p, read_only=True, data_only=False)
    except InvalidFileException:
        raise Erro(400, t("Este formato de planilha não abre aqui. Abra com o programa padrão."))
    except Exception as e:  # corrompida, com senha, zip estragado
        raise Erro(400, t("Não deu para abrir a planilha: {erro}", erro=e))
    abas = []
    try:
        for ws in wb_v.worksheets[:ABAS_PLANILHA]:
            wf = wb_f[ws.title]
            linhas, cortado, largura = [], False, 0
            pares = zip(ws.iter_rows(values_only=True), wf.iter_rows(values_only=True))
            for i, (lv, lf) in enumerate(pares):
                if i >= LINHAS_PLANILHA:
                    cortado = True
                    break
                linha = []
                for j, v in enumerate(lv[:COLUNAS_PLANILHA]):
                    if v is None and j < len(lf) and isinstance(lf[j], str) and lf[j].startswith("="):
                        v = lf[j]  # fórmula que o Excel nunca calculou: mostra a fórmula
                    linha.append(_celula(v))
                if len(lv) > COLUNAS_PLANILHA:
                    cortado = True
                while linha and linha[-1] == "":
                    linha.pop()
                largura = max(largura, len(linha))
                linhas.append(linha)
            while linhas and not linhas[-1]:
                linhas.pop()
            abas.append({"nome": ws.title, "linhas": linhas, "colunas": largura, "cortado": cortado,
                         "total_linhas": ws.max_row, "total_colunas": ws.max_column})
        if len(wb_v.worksheets) > ABAS_PLANILHA:
            abas[-1]["mais_abas"] = len(wb_v.worksheets) - ABAS_PLANILHA
    finally:
        wb_v.close()
        wb_f.close()
    return abas


def _planilha_csv(p: Path) -> list[dict]:
    dados, tamanho = _ler_inicio(p, LIMITE_TEXTO)
    r = decodificar(dados, tamanho > len(dados))
    if r is None:
        raise Erro(400, t("Este arquivo não parece texto, então não dá para mostrar como tabela."))
    texto = r[0]
    if tamanho > len(dados) and "\n" in texto:
        texto = texto[: texto.rfind("\n")]
    amostra = texto[:64 * 1024]
    if p.suffix.lower() == ".tsv":
        separador = "\t"
    else:
        try:
            separador = csv.Sniffer().sniff(amostra, delimiters=",;\t|").delimiter
        except csv.Error:
            contagem = {s: amostra.count(s) for s in (";", ",", "\t", "|")}
            separador = max(contagem, key=contagem.get) if any(contagem.values()) else ","
    linhas, cortado, largura = [], tamanho > len(dados), 0
    for i, linha in enumerate(csv.reader(io.StringIO(texto, newline=""), delimiter=separador)):
        if i >= LINHAS_PLANILHA:
            cortado = True
            break
        if len(linha) > COLUNAS_PLANILHA:
            cortado = True
        linha = linha[:COLUNAS_PLANILHA]
        largura = max(largura, len(linha))
        linhas.append(linha)
    return [{"nome": p.name, "linhas": linhas, "colunas": largura, "cortado": cortado, "separador": separador}]


@rota("GET", "/api/arquivos/planilha")
def r_planilha(h):
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    p = _local(raiz, q.get("caminho"))
    if not p.is_file():
        raise Erro(400, t("{caminho} não é um arquivo.", caminho=_rel(raiz, p)))
    if not _exigir_liberado(h, raiz, p, q):
        return
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        abas, tipo = _planilha_xlsx(p), "xlsx"
    elif ext in (".csv", ".tsv"):
        abas, tipo = _planilha_csv(p), "csv"
    elif ext == ".xls":
        raise Erro(400, t("Planilha no formato antigo (.xls): abra com o programa padrão para ver."))
    else:
        raise Erro(400, t("Este arquivo não é uma planilha (.xlsx ou .csv)."))
    h._json({"caminho": _rel(raiz, p), "tipo": tipo, "abas": abas, "modificado": p.stat().st_mtime,
             "linhas_max": LINHAS_PLANILHA, "colunas_max": COLUNAS_PLANILHA})


@rota("GET", "/api/arquivos/docx")
def r_docx(h):
    from ferramentas import word
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    p = _local(raiz, q.get("caminho"))
    if not p.is_file():
        raise Erro(400, t("{caminho} não é um arquivo.", caminho=_rel(raiz, p)))
    if not _exigir_liberado(h, raiz, p, q):
        return
    try:
        texto = word.ler(p)
    except word.ErroWord as e:
        raise Erro(400, str(e))
    except Exception as e:  # XML estragado
        raise Erro(400, t("Não deu para ler o documento: {erro}", erro=e))
    h._json({"caminho": _rel(raiz, p), "texto": texto, "modificado": p.stat().st_mtime})


def _execucoes_da_conversa(conversa_id: str) -> list[dict]:
    """Os checkpoints da conversa, sem o conteúdo: cada um guarda o arquivo inteiro de antes da mudança, e a tela
    confere a lista a cada resposta. O conteúdo vem depois, só do arquivo pedido (_conteudo_checkpoint)."""
    if not conversa_id:
        return []
    return [dict(r) for r in banco.conexao().execute(
        "SELECT id, caminho, existia FROM checkpoints WHERE conversa_id=? ORDER BY id", (conversa_id,))]


def _tamanho_checkpoint(ck_id: int) -> int:
    # length() de um BLOB vem do cabeçalho do registro: o SQLite não lê o conteúdo para responder
    r = banco.conexao().execute("SELECT length(conteudo) FROM checkpoints WHERE id=?", (ck_id,)).fetchone()
    return int(r[0] or 0) if r else 0


def _conteudo_checkpoint(ck_id: int) -> bytes:
    r = banco.conexao().execute("SELECT conteudo FROM checkpoints WHERE id=?", (ck_id,)).fetchone()
    return bytes(r[0] or b"") if r else b""


def _pedacos_checkpoint(ck_id: int, tamanho: int, bloco: int = 1024 * 1024):
    """O conteúdo do checkpoint em pedaços, sem carregar tudo (blobopen, Python 3.11+; antes disso, lê inteiro)."""
    if tamanho <= 0:
        return
    con = banco.conexao()
    try:
        blob = con.blobopen("checkpoints", "conteudo", ck_id, readonly=True)
    except (AttributeError, sqlite3.Error):
        dados = _conteudo_checkpoint(ck_id)
        for i in range(0, len(dados), bloco):
            yield dados[i:i + bloco]
        return
    with blob:
        while True:
            pedaco = blob.read(bloco)
            if not pedaco:
                break
            yield pedaco


def _mesmo_conteudo(f, pedacos) -> bool:
    """O resto do arquivo aberto `f` é igual aos `pedacos`? Para no primeiro pedaço diferente."""
    for pedaco in pedacos:
        if f.read(len(pedaco)) != pedaco:
            return False
    return f.read(1) == b""


def _mesmo(a: str, b: str) -> bool:
    return os.path.normcase(os.path.normpath(a)) == os.path.normcase(os.path.normpath(b))


def _primeiro_checkpoint(conversa_id: str, p: Path) -> dict | None:
    try:
        alvo = str(p.resolve())
    except (OSError, RuntimeError):
        alvo = str(p)
    for ck in _execucoes_da_conversa(conversa_id):
        if _mesmo(ck["caminho"], alvo) or _mesmo(ck["caminho"], str(p)):
            return ck
    return None


@rota("GET", "/api/arquivos/info")
def r_info(h):
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    p = _local(raiz, q.get("caminho"))
    try:
        st = p.stat()
    except OSError as e:
        raise Erro(400, t("Não deu para ler o arquivo: {erro}", erro=e.strerror or e))
    pasta = p.is_dir()
    d = {"caminho": _rel(raiz, p) if p != raiz else "", "nome": p.name, "absoluto": str(p),
         "tipo": "pasta" if pasta else "arquivo", "modificado": st.st_mtime,
         "criado": getattr(st, "st_birthtime", st.st_ctime), "protegido": _protegido(raiz, p),
         "link": _e_link(p), "executavel": (not pasta) and _executavel(p)}
    if not pasta:
        d.update({"tamanho": st.st_size, "mime": _mime(p), "extensao": p.suffix.lower(), "linguagem": linguagem(p)})
    if q.get("conversa") and not pasta:
        d["mudou"] = _primeiro_checkpoint(str(q["conversa"]), p) is not None
    h._json(d)


@rota("POST", "/api/arquivos/mtimes")
def r_mtimes(h):
    """A data de mudança de várias pastas e arquivos de uma vez (a árvore e o visualizador conferem a cada poucos
    segundos se algo mudou, sem listar tudo de novo)."""
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    saida = {}
    for c in list(d.get("caminhos") or [])[:500]:
        try:
            p = _local(raiz, c)
            saida[str(c)] = p.stat().st_mtime
        except (Erro, OSError):
            saida[str(c)] = None
    h._json({"mtimes": saida})


@rota("GET", "/api/arquivos/mudancas")
def r_mudancas(h):
    """O que o Claude mudou nesta conversa. Sem `caminho`: a lista dos arquivos (dentro da raiz). Com `caminho`: o
    conteúdo de antes da primeira mudança, o de agora e o diff entre os dois."""
    q = h.query()
    raiz = _raiz(q.get("raiz"))
    conversa = str(q.get("conversa") or "")
    if not conversa:
        raise Erro(400, t("Falta dizer qual é a conversa."))
    if not q.get("caminho"):
        vistos, lista = set(), []
        for ck in _execucoes_da_conversa(conversa):
            p = Path(ck["caminho"])
            chave = os.path.normcase(str(p))
            if chave in vistos:
                continue
            vistos.add(chave)
            try:
                if p.resolve().is_relative_to(raiz):
                    lista.append(_rel(raiz, p))
            except (OSError, RuntimeError, ValueError):
                continue
        return h._json({"arquivos": lista})
    p = _local(raiz, q.get("caminho"), existir=False)
    ck = _primeiro_checkpoint(conversa, p)
    if not ck:
        return h._json({"mudou": False, "caminho": _rel(raiz, p)})
    if not _exigir_liberado(h, raiz, p, q):
        return
    existia = bool(ck["existia"])
    tam_antes = _tamanho_checkpoint(ck["id"]) if existia else 0
    grande = tam_antes > LIMITE_TEXTO
    depois, tam_depois, igual = b"", 0, None
    existe = p.is_file()
    if existe:
        # lê só o necessário: o começo para comparar na tela; o arquivo inteiro só quando o tamanho é igual ao de antes
        # e ele é grande demais para a tela (aí compara em pedaços e para na primeira diferença)
        try:
            with open(p, "rb") as f:
                tam_depois = os.fstat(f.fileno()).st_size
                grande = grande or tam_depois > LIMITE_TEXTO
                if tam_depois != tam_antes:
                    igual = False
                elif grande:
                    igual = _mesmo_conteudo(f, _pedacos_checkpoint(ck["id"], tam_antes))
                if not grande:
                    depois = f.read(LIMITE_TEXTO + 1)
                    grande = len(depois) > LIMITE_TEXTO  # cresceu enquanto lia
        except OSError as e:
            raise Erro(400, t("Não deu para ler o arquivo: {erro}", erro=e.strerror or e))
    antes = _conteudo_checkpoint(ck["id"]) if existia and not grande else b""
    if igual is None:
        igual = antes == depois
    base = {"mudou": True, "caminho": _rel(raiz, p), "criado": not existia, "apagado": not existe,
            "igual": bool(igual) and existe == existia, "tamanho_antes": tam_antes, "tamanho_depois": tam_depois}
    if grande:
        return h._json({**base, "grande": True})
    ta, td = decodificar(antes), decodificar(depois)
    if ta is None or td is None:
        return h._json({**base, "binario": True})
    rel = _rel(raiz, p)
    linhas = difflib.unified_diff(ta[0].splitlines(), td[0].splitlines(), f"a/{rel}", f"b/{rel}", lineterm="", n=3)
    diff = list(linhas)
    h._json({**base, "binario": False, "antes": ta[0], "depois": td[0], "diff": "\n".join(diff[:4000]),
             "diff_cortado": len(diff) > 4000})


# ======================= rotas que mudam algo =======================

@rota("POST", "/api/arquivos/criar")
def r_criar(h):
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    pasta = _local(raiz, d.get("pasta") or "")
    if not pasta.is_dir():
        raise Erro(400, t("{caminho} não é uma pasta.", caminho=_rel(raiz, pasta)))
    nome = _validar_nome(d.get("nome"))
    alvo = pasta / nome
    if alvo.exists() or _e_link(alvo):
        raise Erro(409, t("Já existe um item chamado {nome} aqui.", nome=nome))
    try:
        if d.get("tipo") == "pasta":
            alvo.mkdir()
        else:
            with open(alvo, "x", encoding="utf-8"):
                pass
    except FileExistsError:
        raise Erro(409, t("Já existe um item chamado {nome} aqui.", nome=nome))
    except OSError as e:
        raise Erro(400, t("Não deu para criar: {erro}", erro=e.strerror or e))
    h._json({"ok": True, "caminho": _rel(raiz, alvo), "tipo": "pasta" if d.get("tipo") == "pasta" else "arquivo"})


@rota("POST", "/api/arquivos/renomear")
def r_renomear(h):
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    p = _local(raiz, d.get("caminho"), seguir=False)
    if p == raiz:
        raise Erro(400, t("A pasta do trabalho não pode ser renomeada por aqui."))
    nome = _validar_nome(d.get("nome"))
    novo = p.with_name(nome)
    if nome == p.name:
        return h._json({"ok": True, "caminho": _rel(raiz, p)})
    so_maiusculas = nome.casefold() == p.name.casefold()
    if (novo.exists() or _e_link(novo)) and not so_maiusculas:
        raise Erro(409, t("Já existe um item chamado {nome} aqui.", nome=nome))
    try:
        os.rename(p, novo)
    except PermissionError:
        raise Erro(409, t("O Windows não deixou renomear: o arquivo pode estar aberto em outro programa."))
    except OSError as e:
        raise Erro(400, t("Não deu para renomear: {erro}", erro=e.strerror or e))
    h._json({"ok": True, "caminho": _rel(raiz, novo)})


# ---------- Lixeira do Windows (SHFileOperationW) ----------

FO_DELETE = 0x0003
FOF_SILENT = 0x0004
FOF_NOCONFIRMATION = 0x0010
FOF_ALLOWUNDO = 0x0040
FOF_NOERRORUI = 0x0400
FOF_WANTNUKEWARNING = 0x4000  # se o Windows não puder pôr na Lixeira (grande demais), pergunta antes de apagar de vez
DRIVE_FIXED = 3

if _no_windows():
    from ctypes import wintypes

    class SHFILEOPSTRUCTW(ctypes.Structure):
        if ctypes.sizeof(ctypes.c_void_p) == 4:
            _pack_ = 1  # no Windows de 32 bits a estrutura é empacotada byte a byte (shellapi.h)
        _fields_ = [("hwnd", wintypes.HWND), ("wFunc", wintypes.UINT), ("pFrom", wintypes.LPCWSTR),
                    ("pTo", wintypes.LPCWSTR), ("fFlags", ctypes.c_ushort), ("fAnyOperationsAborted", wintypes.BOOL),
                    ("hNameMappings", ctypes.c_void_p), ("lpszProgressTitle", wintypes.LPCWSTR)]

    # instâncias próprias das DLLs: os argtypes daqui não vazam para outros módulos (nem os deles para cá)
    _k32 = ctypes.WinDLL("kernel32")
    _k32.GetDriveTypeW.argtypes = [wintypes.LPCWSTR]
    _k32.GetDriveTypeW.restype = wintypes.UINT
    _shell32 = ctypes.WinDLL("shell32")
    _shell32.SHFileOperationW.argtypes = [ctypes.POINTER(SHFILEOPSTRUCTW)]
    _shell32.SHFileOperationW.restype = ctypes.c_int
    _user32 = ctypes.WinDLL("user32")
    _user32.GetForegroundWindow.argtypes = []
    _user32.GetForegroundWindow.restype = wintypes.HWND


def _janela_dona():
    """A janela da frente (a do Clave, que acabou de pedir a Lixeira): a pergunta "apagar de vez?" do Windows fica
    presa a ela e aparece por cima. Sem dona, ela podia nascer atrás do Clave e o pedido ficava parado esperando."""
    if not _no_windows():
        return None
    try:
        return _user32.GetForegroundWindow() or None
    except OSError:
        return None


def _tem_lixeira(p: Path) -> bool:
    """A Lixeira só existe nos discos do próprio computador: num pendrive ou na rede, o Windows apagaria de vez."""
    if not _no_windows():
        return True
    raiz = os.path.splitdrive(str(p))[0]
    if not raiz or raiz.startswith("\\\\"):
        return False
    return _k32.GetDriveTypeW(raiz + "\\") == DRIVE_FIXED


def mandar_para_lixeira(p: Path) -> None:
    """Põe o arquivo ou a pasta na Lixeira do Windows. Levanta OSError se não deu."""
    if not _no_windows():
        raise OSError(t("A Lixeira só funciona no Windows."))
    # a lista de caminhos termina com dois nulos: o buffer guarda o caminho, um nulo nosso e o nulo final dele
    origem = ctypes.create_unicode_buffer(str(p) + "\0")
    op = SHFILEOPSTRUCTW()
    op.hwnd = _janela_dona()
    op.wFunc = FO_DELETE
    op.pFrom = ctypes.cast(origem, wintypes.LPCWSTR)
    op.pTo = None
    op.fFlags = FOF_ALLOWUNDO | FOF_NOCONFIRMATION | FOF_SILENT | FOF_NOERRORUI | FOF_WANTNUKEWARNING
    r = _shell32.SHFileOperationW(ctypes.byref(op))
    if r != 0:
        raise OSError(t("o Windows recusou (código {codigo})", codigo=hex(r)))
    if op.fAnyOperationsAborted:
        raise OSError(t("a operação foi cancelada"))


@rota("POST", "/api/arquivos/lixeira")
def r_lixeira(h):
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    p = _local(raiz, d.get("caminho"), seguir=False)
    if p == raiz:
        raise Erro(400, t("A pasta do trabalho inteira não vai para a Lixeira por aqui."))
    if _e_link(p):
        raise Erro(400, t("Isto é um atalho de pasta (link). Para tirar só o atalho, use o Explorador do Windows."))
    if not _tem_lixeira(p):
        raise Erro(400, t("Esta pasta fica num pendrive ou na rede, onde o Windows não tem Lixeira. Apague pelo Explorador se tiver certeza."))
    try:
        mandar_para_lixeira(p)
    except OSError as e:
        raise Erro(409, t("Não deu para mandar para a Lixeira: {erro}. Ele pode estar aberto em outro programa.",
                          erro=e.strerror or e))
    if p.exists():
        raise Erro(409, t("O Windows não tirou o item do lugar. Ele pode estar aberto em outro programa."))
    h._json({"ok": True, "caminho": _rel(raiz, p)})


@rota("POST", "/api/arquivos/abrir")
def r_abrir(h):
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    p = _local(raiz, d.get("caminho"))
    if p.is_file() and _executavel(p) and not d.get("confirmado"):
        return h._json({"erro": t("Este arquivo é um programa ou script: abrir vai executá-lo."),
                        "codigo": "executavel"}, 409)
    try:
        if _no_windows():
            os.startfile(str(p))
        else:
            subprocess.Popen(["xdg-open", str(p)])
    except OSError as e:
        raise Erro(400, t("Não deu para abrir: {erro}", erro=e.strerror or e))
    h._json({"ok": True})


@rota("POST", "/api/arquivos/revelar")
def r_revelar(h):
    d = h.corpo()
    raiz = _raiz(d.get("raiz"))
    p = _local(raiz, d.get("caminho"), seguir=False)
    try:
        if _no_windows():
            # a vírgula colada e as aspas em volta do caminho são o formato que o Explorador entende (aspas não
            # existem em nome de arquivo do Windows, então o caminho não fecha a aspa)
            alvo = f'explorer.exe /select,"{p}"' if p != raiz else f'explorer.exe "{p}"'
            subprocess.Popen(alvo, creationflags=0x08000000)
        else:
            subprocess.Popen(["xdg-open", str(p.parent if p.is_file() else p)])
    except OSError as e:
        raise Erro(400, t("Não deu para abrir o Explorador: {erro}", erro=e.strerror or e))
    h._json({"ok": True})
