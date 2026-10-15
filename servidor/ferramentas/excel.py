"""Planilhas do Excel: ler, criar e alterar (.xlsx/.xlsm com openpyxl) e usar o próprio Excel quando ele existe.

O openpyxl escreve fórmulas mas não calcula: o valor aparece quando o arquivo é
aberto no Excel. Para conferir resultado de fórmula sem abrir, `no_excel` usa o
Excel instalado (pelo PowerShell) para recalcular e salvar, exportar PDF ou
converter .xls e .csv em .xlsx.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from i18n import t

EXTENSOES = {".xlsx", ".xlsm"}
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ErroExcel(Exception):
    pass


# ---------------- leitura ----------------

def _valor_txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat(sep=" ") if isinstance(v, dt.datetime) and (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return f"{v:.10g}"
    return str(v).replace("\n", "\\n")


def ler(caminho: str | Path, aba: str | None = None, intervalo: str | None = None, max_linhas: int = 100,
        max_colunas: int = 30) -> str:
    p = Path(caminho)
    if p.suffix.lower() not in EXTENSOES:
        raise ErroExcel(t("Só leio .xlsx e .xlsm. Para .xls ou .csv, converta antes (no_excel com acao converter_xlsx) ou leia o .csv como texto."))
    if not p.is_file():
        raise ErroExcel(t("Arquivo não encontrado: {caminho}", caminho=p))
    grande = p.stat().st_size > 8 * 1024 * 1024
    try:
        wb_f = openpyxl.load_workbook(p, data_only=False, read_only=grande)
        wb_v = openpyxl.load_workbook(p, data_only=True, read_only=grande)
    except Exception as e:  # arquivo corrompido, protegido por senha etc.
        raise ErroExcel(t("Não deu para abrir a planilha: {erro}", erro=e))
    try:
        linhas = [t("Planilha: {nome}", nome=p.name), t("Abas: {abas}", abas=", ".join(
            t("{aba} ({linhas} linhas x {colunas} colunas)", aba=ws.title, linhas=ws.max_row, colunas=ws.max_column)
            for ws in wb_f.worksheets))]
        alvos = [wb_f[aba]] if aba else wb_f.worksheets[:1] if intervalo else wb_f.worksheets[:3]
        if aba and aba not in wb_f.sheetnames:
            raise ErroExcel(t("Não existe a aba '{aba}'. Abas: {abas}", aba=aba, abas=", ".join(wb_f.sheetnames)))
        sem_valor = False
        for ws in alvos:
            wv = wb_v[ws.title]
            if intervalo:
                c1, l1, c2, l2 = range_boundaries(intervalo.upper())
            else:
                c1, l1, c2, l2 = 1, 1, ws.max_column or 1, ws.max_row or 1
                if grande and (not ws.max_row or ws.max_row < 2):
                    l2 = l1 + max_linhas - 1
            c2 = min(c2, c1 + max_colunas - 1)
            l2_corte = min(l2, l1 + max_linhas - 1)
            linhas.append("")
            linhas.append(t("## Aba \"{aba}\" ({intervalo})", aba=ws.title,
                            intervalo=f"{get_column_letter(c1)}{l1}:{get_column_letter(c2)}{l2_corte}"))
            if not grande:
                extras = []
                if ws.freeze_panes:
                    extras.append(t("painel congelado em {celula}", celula=ws.freeze_panes))
                if ws.tables:
                    extras.append(t("tabelas: {lista}", lista=", ".join(f"{n} ({ref})" for n, ref in ws.tables.items())))
                if ws.merged_cells.ranges:
                    extras.append(t("{n} mesclagem(ns)", n=len(ws.merged_cells.ranges)))
                if getattr(ws, "_charts", None):
                    extras.append(t("{n} gráfico(s)", n=len(ws._charts)))
                if extras:
                    linhas.append("(" + "; ".join(extras) + ")")
            linhas.append(t("Cada linha: número | coluna=valor (fórmulas aparecem como =FÓRMULA → valor calculado)."))
            vazias = 0
            faixa = dict(min_row=l1, max_row=l2_corte, min_col=c1, max_col=c2, values_only=True)
            for r, (lf, lv) in enumerate(zip(ws.iter_rows(**faixa), wv.iter_rows(**faixa)), start=l1):
                partes = []
                for j, vf in enumerate(lf):
                    if vf is None:
                        continue
                    letra = get_column_letter(c1 + j)
                    if isinstance(vf, str) and vf.startswith("="):
                        vv = lv[j] if j < len(lv) else None
                        if vv is None:
                            sem_valor = True
                        partes.append(f"{letra}={vf} → {_valor_txt(vv) if vv is not None else '?'}")
                    else:
                        partes.append(f"{letra}={_valor_txt(vf)}")
                if partes:
                    linhas.append(f"{r} | " + " | ".join(partes))
                    vazias = 0
                else:
                    vazias += 1
                    if vazias == 1:
                        linhas.append(f"{r} | " + t("(vazia)"))
            if l2 > l2_corte:
                linhas.append(t("[... mais {n} linhas; leia com intervalo, ex.: {exemplo}]", n=l2 - l2_corte,
                                exemplo=f"{get_column_letter(c1)}{l2_corte + 1}:{get_column_letter(c2)}"
                                        f"{min(l2, l2_corte + max_linhas)}"))
        if len(wb_f.worksheets) > len(alvos) and not aba:
            linhas.append("\n" + t("[mostrando {n} de {total} abas; use o parâmetro aba para ler as outras]",
                                   n=len(alvos), total=len(wb_f.worksheets)))
        if sem_valor:
            linhas.append("\n" + t("'?' = fórmula sem valor guardado (o arquivo nunca foi recalculado no Excel). Para ver o resultado, use no_excel com acao recalcular, ou calcule você mesmo."))
        return "\n".join(linhas)
    finally:
        wb_f.close()
        wb_v.close()


# ---------------- escrita ----------------

_DATA = re.compile(r"^\d{4}-\d{2}-\d{2}( \d{2}:\d{2}(:\d{2})?)?$")
_CORES = {"vermelho": "FF0000", "verde": "00B050", "azul": "0070C0", "amarelo": "FFFF00", "laranja": "FFC000",
          "cinza": "BFBFBF", "preto": "000000", "branco": "FFFFFF", "roxo": "7030A0", "verde_claro": "C6EFCE",
          "vermelho_claro": "FFC7CE", "amarelo_claro": "FFEB9C", "azul_claro": "DDEBF7"}


def _cor(c) -> str | None:
    if not c:
        return None
    c = str(c).strip().lower()
    c = _CORES.get(c, c).lstrip("#").upper()
    if re.fullmatch(r"[0-9A-F]{6}", c):
        return c
    raise ErroExcel(t("Cor inválida: {cor} (use #RRGGBB ou nomes como verde, vermelho_claro)", cor=c))


def _converter(v, datas: bool):
    if isinstance(v, str) and datas and _DATA.match(v):
        try:
            return dt.datetime.fromisoformat(v) if " " in v else dt.datetime.strptime(v, "%Y-%m-%d")
        except ValueError:
            return v
    return v


def _aba(wb, nome: str | None, criar=False):
    if not nome:
        return wb.active
    if nome in wb.sheetnames:
        return wb[nome]
    if criar:
        return wb.create_sheet(nome[:31])
    raise ErroExcel(t("Não existe a aba '{aba}'. Abas: {abas}", aba=nome, abas=", ".join(wb.sheetnames)))


def _celulas(ws, intervalo: str):
    c1, l1, c2, l2 = range_boundaries(intervalo.upper())
    for r in range(l1, l2 + 1):
        for c in range(c1, c2 + 1):
            yield ws.cell(row=r, column=c)


def _op(wb, op: dict, datas: bool) -> str:
    acao = op.get("acao")
    if acao == "criar_aba":
        nome = str(op["aba"])[:31]
        if nome in wb.sheetnames:
            return t("aba '{aba}' já existia", aba=nome)
        wb.create_sheet(nome, op.get("posicao"))
        return t("aba '{aba}' criada", aba=nome)
    if acao == "renomear_aba":
        ws = _aba(wb, op["aba"])
        ws.title = str(op["novo_nome"])[:31]
        return t("aba renomeada para '{aba}'", aba=ws.title)
    if acao == "apagar_aba":
        ws = _aba(wb, op["aba"])
        if len(wb.worksheets) == 1:
            raise ErroExcel(t("Não dá para apagar a única aba."))
        wb.remove(ws)
        return t("aba '{aba}' apagada", aba=op["aba"])
    ws = _aba(wb, op.get("aba"), criar=acao in ("escrever", "celulas"))
    if acao == "escrever":
        linhas = op.get("linhas") or []
        c0, l0 = _coluna_linha(op.get("inicio") or "A1")
        for i, linha in enumerate(linhas):
            for j, v in enumerate(linha if isinstance(linha, list) else [linha]):
                ws.cell(row=l0 + i, column=c0 + j, value=_converter(v, datas))
        return t("{n} linha(s) escritas em '{aba}' a partir de {inicio}", n=len(linhas), aba=ws.title,
                 inicio=op.get("inicio") or "A1")
    if acao == "celulas":
        valores = op.get("valores") or {}
        for ref, v in valores.items():
            ws[ref.upper()] = _converter(v, datas)
        return t("{n} célula(s) definidas em '{aba}'", n=len(valores), aba=ws.title)
    if acao == "limpar":
        n = 0
        for cel in _celulas(ws, op["intervalo"]):
            cel.value = None
            n += 1
        return t("{n} célula(s) limpas", n=n)
    if acao == "formatar":
        n = 0
        for cel in _celulas(ws, op["intervalo"]):
            f = cel.font
            if any(k in op for k in ("negrito", "italico", "cor_texto", "tamanho_fonte", "sublinhado")):
                cel.font = Font(name=f.name, size=op.get("tamanho_fonte", f.size), bold=op.get("negrito", f.bold),
                                italic=op.get("italico", f.italic), underline="single" if op.get("sublinhado") else f.underline,
                                color=_cor(op.get("cor_texto")) or f.color)
            if op.get("cor_fundo"):
                cel.fill = PatternFill("solid", fgColor=_cor(op["cor_fundo"]))
            if op.get("formato_numero"):
                cel.number_format = str(op["formato_numero"])
            if op.get("alinhamento") or op.get("quebra_texto") is not None:
                h = {"esquerda": "left", "centro": "center", "direita": "right"}.get(op.get("alinhamento"), cel.alignment.horizontal)
                cel.alignment = Alignment(horizontal=h, vertical="center" if op.get("alinhamento") else cel.alignment.vertical,
                                          wrap_text=op.get("quebra_texto", cel.alignment.wrap_text))
            if op.get("borda"):
                lado = Side(style={"fina": "thin", "grossa": "medium", "nenhuma": None}.get(op["borda"], "thin"))
                cel.border = Border(left=lado, right=lado, top=lado, bottom=lado)
            n += 1
        return t("{n} célula(s) formatadas", n=n)
    if acao == "largura_colunas":
        if op.get("auto"):
            larguras = {}
            for linha in ws.iter_rows():
                for cel in linha:
                    if cel.value is not None:
                        letra = cel.column_letter
                        larguras[letra] = max(larguras.get(letra, 0), min(60, len(_valor_txt(cel.value)) + 2))
            for letra, w in larguras.items():
                ws.column_dimensions[letra].width = max(8, w)
            return t("largura automática em {n} coluna(s)", n=len(larguras))
        for letra, w in (op.get("larguras") or {}).items():
            ws.column_dimensions[letra.upper()].width = float(w)
        return t("larguras ajustadas")
    if acao == "congelar":
        ws.freeze_panes = op.get("celula") or None
        return t("painel congelado em {celula}", celula=op.get("celula"))
    if acao == "filtro":
        ws.auto_filter.ref = op["intervalo"].upper()
        return t("filtro em {intervalo}", intervalo=op["intervalo"])
    if acao == "mesclar":
        ws.merge_cells(op["intervalo"].upper())
        return t("células {intervalo} mescladas", intervalo=op["intervalo"])
    if acao == "tabela":
        nome = re.sub(r"\W", "_", str(op.get("nome") or f"Tabela{len(ws.tables) + 1}"))
        tab = Table(displayName=nome, ref=op["intervalo"].upper())
        tab.tableStyleInfo = TableStyleInfo(name=op.get("estilo") or "TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
        return t("tabela {nome} criada em {intervalo}", nome=nome, intervalo=op["intervalo"])
    if acao in ("inserir_linhas", "apagar_linhas"):
        fn = ws.insert_rows if acao == "inserir_linhas" else ws.delete_rows
        fn(int(op["linha"]), int(op.get("quantidade") or 1))
        return t("{acao}: {n} a partir da linha {linha}", acao=t(acao.replace("_", " ")), n=op.get("quantidade") or 1,
                 linha=op["linha"])
    if acao in ("inserir_colunas", "apagar_colunas"):
        fn = ws.insert_cols if acao == "inserir_colunas" else ws.delete_cols
        fn(column_index_from_string(str(op["coluna"]).upper()), int(op.get("quantidade") or 1))
        return t("{acao}: {n} a partir da coluna {coluna}", acao=t(acao.replace("_", " ")), n=op.get("quantidade") or 1,
                 coluna=op["coluna"])
    if acao == "validacao_lista":
        opcoes = [str(o).replace('"', "'") for o in op.get("opcoes") or []]
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(op["intervalo"].upper())
        return t("lista de opções em {intervalo}", intervalo=op["intervalo"])
    if acao == "formatacao_condicional":
        regra = op.get("regra")
        intervalo = op["intervalo"].upper()
        if regra == "escala_cores":
            ws.conditional_formatting.add(intervalo, ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                                                                    mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
        elif regra == "barra_dados":
            ws.conditional_formatting.add(intervalo, DataBarRule(start_type="min", end_type="max", color=_cor(op.get("cor_fundo")) or "638EC6"))
        else:
            operadores = {"maior_que": "greaterThan", "menor_que": "lessThan", "igual": "equal", "entre": "between",
                          "maior_ou_igual": "greaterThanOrEqual", "menor_ou_igual": "lessThanOrEqual", "diferente": "notEqual"}
            if regra not in operadores:
                raise ErroExcel(t("Regra desconhecida: {regra}", regra=regra))
            formula = [str(op.get("valor"))] + ([str(op.get("valor2"))] if regra == "entre" else [])
            ws.conditional_formatting.add(intervalo, CellIsRule(operator=operadores[regra], formula=formula,
                                                                fill=PatternFill("solid", fgColor=_cor(op.get("cor_fundo") or "amarelo_claro")),
                                                                font=Font(color=_cor(op.get("cor_texto"))) if op.get("cor_texto") else None))
        return t("formatação condicional ({regra}) em {intervalo}", regra=regra, intervalo=intervalo)
    if acao == "grafico":
        return _grafico(ws, op)
    raise ErroExcel(t("Ação desconhecida: {acao}", acao=acao))


def _coluna_linha(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Za-z]{1,3})(\d+)", ref.strip())
    if not m:
        raise ErroExcel(t("Célula inválida: {celula}", celula=ref))
    return column_index_from_string(m.group(1).upper()), int(m.group(2))


def _grafico(ws, op: dict) -> str:
    tipo = op.get("tipo", "coluna")
    classes = {"coluna": BarChart, "barra": BarChart, "linha": LineChart, "pizza": PieChart, "area": AreaChart,
               "dispersao": ScatterChart}
    if tipo not in classes:
        raise ErroExcel(t("Tipo de gráfico desconhecido: {tipo}", tipo=tipo))
    g = classes[tipo]()
    if tipo == "barra":
        g.type = "bar"
    elif tipo == "coluna":
        g.type = "col"
    c1, l1, c2, l2 = range_boundaries(op["dados"].upper())
    dados = Reference(ws, min_col=c1, min_row=l1, max_col=c2, max_row=l2)
    cats = None
    if op.get("categorias"):
        a1, b1, a2, b2 = range_boundaries(op["categorias"].upper())
        cats = Reference(ws, min_col=a1, min_row=b1, max_col=a2, max_row=b2)
    if tipo == "dispersao":
        if not cats:
            raise ErroExcel(t("Gráfico de dispersão precisa de 'categorias' (os valores do eixo X)."))
        for c in range(c1, c2 + 1):
            ys = Reference(ws, min_col=c, min_row=l1 + 1, max_row=l2)
            g.series.append(Series(ys, cats, title_from_data=False, title=str(ws.cell(row=l1, column=c).value or "")))
    else:
        g.add_data(dados, titles_from_data=bool(op.get("com_cabecalho", True)))
        if cats:
            g.set_categories(cats)
    if op.get("titulo"):
        g.title = op["titulo"]
    if tipo != "pizza":
        if op.get("titulo_x"):
            g.x_axis.title = op["titulo_x"]
        if op.get("titulo_y"):
            g.y_axis.title = op["titulo_y"]
    g.width = float(op.get("largura") or 18)
    g.height = float(op.get("altura") or 9)
    ws.add_chart(g, (op.get("posicao") or "H2").upper())
    return t("gráfico de {tipo} em {posicao}", tipo=t(tipo), posicao=op.get("posicao") or "H2")


def aplicar(caminho: str | Path, operacoes: list, criar: bool = True, converter_datas: bool = True,
            base: str | Path | None = None) -> str:
    """Cria ou altera uma planilha. `base` = começar de uma cópia deste arquivo."""
    p = Path(caminho)
    if p.suffix.lower() not in EXTENSOES:
        raise ErroExcel(t("O arquivo precisa terminar em .xlsx (ou .xlsm)."))
    if not isinstance(operacoes, list) or not operacoes:
        raise ErroExcel(t("Mande pelo menos uma operação."))
    origem = Path(base) if base else p
    if origem.is_file():
        try:
            wb = openpyxl.load_workbook(origem, keep_vba=origem.suffix.lower() == ".xlsm")
        except Exception as e:
            raise ErroExcel(t("Não deu para abrir a planilha: {erro}", erro=e))
        novo = False
    elif criar:
        if p.suffix.lower() == ".xlsm":
            raise ErroExcel(t("Planilha nova precisa ser .xlsx (.xlsm é só para arquivo que já tem macro)."))
        wb = openpyxl.Workbook()
        primeira = next((o.get("aba") for o in operacoes if isinstance(o, dict) and o.get("aba")), None)
        if primeira:
            wb.active.title = str(primeira)[:31]
        novo = True
    else:
        raise ErroExcel(t("Arquivo não encontrado: {caminho}", caminho=p))
    feitos = []
    for i, op in enumerate(operacoes, 1):
        if not isinstance(op, dict):
            raise ErroExcel(t("Operação {i} inválida.", i=i))
        try:
            feitos.append(_op(wb, op, converter_datas))
        except ErroExcel as e:
            raise ErroExcel(t("Operação {i} ({acao}): {erro}", i=i, acao=op.get("acao"), erro=e))
        except (KeyError, ValueError, TypeError) as e:
            raise ErroExcel(t("Operação {i} ({acao}): parâmetro faltando ou inválido ({erro})", i=i, acao=op.get("acao"),
                              erro=e))
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_name(f".~clave-{p.name}")
    wb.save(tmp)
    wb.close()
    try:
        os.replace(tmp, p)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise ErroExcel(t("Não deu para salvar: o arquivo está aberto no Excel. Feche ele e tente de novo."))
    cab = t("Planilha criada: {caminho}", caminho=p) if novo else t("Planilha alterada: {caminho}", caminho=p)
    return cab + "\n- " + "\n- ".join(feitos) + ("\n" + t("(Fórmulas são calculadas quando o arquivo abre no Excel.)")
                                                 if any(_tem_formula(o) for o in operacoes) else "")


def _tem_formula(op) -> bool:
    texto = json.dumps(op, ensure_ascii=False)
    return '"=' in texto


# ---------------- o próprio Excel (PowerShell + COM) ----------------

def excel_instalado() -> bool:
    if sys.platform != "win32":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Excel.Application"))
        return True
    except OSError:
        return False


_PS = r"""
$ErrorActionPreference = 'Stop'
$entrada = $env:CLAVE_EXCEL_ENTRADA; $saida = $env:CLAVE_EXCEL_SAIDA; $acao = $env:CLAVE_EXCEL_ACAO
$x = New-Object -ComObject Excel.Application
$x.Visible = $false; $x.DisplayAlerts = $false; $x.AskToUpdateLinks = $false
$x.AutomationSecurity = 3   # msoAutomationSecurityForceDisable: nenhuma macro roda
$x.EnableEvents = $false    # nem eventos de abertura
try {
  $wb = $x.Workbooks.Open($entrada, 0, $false)
  $x.CalculateFull()
  if ($acao -eq 'recalcular') { $wb.Save() }
  elseif ($acao -eq 'pdf') { $wb.ExportAsFixedFormat(0, $saida) }
  elseif ($acao -eq 'converter_xlsx') { $wb.SaveAs($saida, 51) }
  $wb.Close($false)
  'ok'
} finally {
  $x.Quit()
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($x)
}
"""


def no_excel(caminho: str | Path, acao: str, saida: str | Path | None = None, tempo: int = 180) -> str:
    if not excel_instalado():
        raise ErroExcel(t("O Microsoft Excel não está instalado neste computador, então não dá para recalcular, exportar PDF ou converter por ele."))
    p = Path(caminho).resolve()
    if not p.is_file():
        raise ErroExcel(t("Arquivo não encontrado: {caminho}", caminho=p))
    if acao not in ("recalcular", "pdf", "converter_xlsx"):
        raise ErroExcel(t("Ação inválida (use recalcular, pdf ou converter_xlsx)."))
    if acao == "pdf":
        saida = Path(saida).resolve() if saida else p.with_suffix(".pdf")
    elif acao == "converter_xlsx":
        saida = Path(saida).resolve() if saida else p.with_suffix(".xlsx")
        if saida == p:
            raise ErroExcel(t("O arquivo já é .xlsx."))
    env = dict(os.environ, CLAVE_EXCEL_ENTRADA=str(p), CLAVE_EXCEL_SAIDA=str(saida or ""), CLAVE_EXCEL_ACAO=acao)
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as f:
        f.write(_PS)
        script = f.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", script],
                           capture_output=True, text=True, timeout=tempo, env=env, creationflags=0x08000000)
    except subprocess.TimeoutExpired:
        raise ErroExcel(t("O Excel demorou demais para responder (ele pode estar esperando alguma janela aberta)."))
    finally:
        os.unlink(script)
    if r.returncode != 0 or "ok" not in r.stdout:
        erro = (r.stderr or r.stdout).strip().splitlines()
        raise ErroExcel(t("O Excel não conseguiu: {erro}", erro=(erro[0] if erro else t("erro desconhecido"))[:300]))
    if acao == "recalcular":
        return t("Fórmulas recalculadas e salvas no próprio Excel: {caminho}. Leia de novo para ver os valores.", caminho=p)
    return t("PDF exportado: {caminho}", caminho=saida) if acao == "pdf" else t("Convertido para .xlsx: {caminho}",
                                                                                caminho=saida)


SEGUROS_PARA_ABRIR = {".xlsx", ".csv", ".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".txt", ".md"}


def abrir_no_excel(caminho: str | Path) -> str:
    p = Path(caminho).resolve()
    if not p.is_file():
        raise ErroExcel(t("Arquivo não encontrado: {caminho}", caminho=p))
    if p.suffix.lower() not in SEGUROS_PARA_ABRIR:
        raise ErroExcel(t("Esse tipo de arquivo não é aberto pelo Clave (ele poderia executar um programa)."))
    if sys.platform == "win32":
        os.startfile(str(p))  # abre no programa padrão (o Excel, quando instalado)
    else:
        subprocess.Popen(["xdg-open" if shutil.which("xdg-open") else "open", str(p)])
    return t("Abri {nome} no Excel (ou no programa padrão de planilhas) para a pessoa ver.", nome=p.name)


# ---------------- esquema das operações (para as ferramentas) ----------------

# a descrição é traduzida pelo catálogo das ferramentas (i18n.traduzir_ferramentas)
ESQUEMA_OPERACOES = {
    "type": "array",
    "description": (
        "Lista de operações, aplicadas em ordem. Cada uma é um objeto com 'acao' e os campos dela:\n"
        "- escrever: aba, inicio (ex. 'A1'), linhas (lista de listas). Texto que começa com '=' vira fórmula "
        "(em inglês, separador vírgula: =SUM(B2:B10), =IF(A2>0,\"sim\",\"não\")). Datas como '2026-10-14' viram data.\n"
        "- celulas: aba, valores ({'A1': 'Nome', 'B5': '=AVERAGE(B2:B4)'})\n"
        "- limpar: aba, intervalo\n"
        "- formatar: aba, intervalo, e qualquer um de: negrito, italico, sublinhado, cor_texto, cor_fundo "
        "(#RRGGBB ou verde, vermelho_claro...), tamanho_fonte, formato_numero (ex. 'R$ #,##0.00', '0.0%', "
        "'dd/mm/yyyy'), alinhamento (esquerda|centro|direita), quebra_texto, borda (fina|grossa|nenhuma)\n"
        "- largura_colunas: aba, larguras ({'A': 30}) ou auto: true\n"
        "- congelar: aba, celula (ex. 'A2' congela a linha 1)\n"
        "- filtro: aba, intervalo   - mesclar: aba, intervalo\n"
        "- tabela: aba, intervalo, nome, estilo (opcional, ex. TableStyleMedium9)\n"
        "- grafico: aba, tipo (coluna|barra|linha|pizza|area|dispersao), dados (intervalo com o cabeçalho na "
        "primeira linha), categorias (rótulos ou eixo X, sem cabeçalho), posicao, titulo, titulo_x, titulo_y\n"
        "- formatacao_condicional: aba, intervalo, regra (maior_que|menor_que|maior_ou_igual|menor_ou_igual|igual|"
        "diferente|entre|escala_cores|barra_dados), valor, valor2, cor_fundo, cor_texto\n"
        "- validacao_lista: aba, intervalo, opcoes (lista de textos para uma caixa de seleção)\n"
        "- inserir_linhas / apagar_linhas: aba, linha, quantidade   - inserir_colunas / apagar_colunas: aba, coluna, quantidade\n"
        "- criar_aba: aba, posicao   - renomear_aba: aba, novo_nome   - apagar_aba: aba"),
    "items": {"type": "object", "properties": {"acao": {"type": "string"}}, "required": ["acao"]},
}
