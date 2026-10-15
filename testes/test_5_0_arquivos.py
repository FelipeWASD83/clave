"""5.0: as abas Arquivos e Visualizador (servidor/explorador.py): lista, busca, conteúdo, bruto com Range, planilhas,
docx, protegidos, criar, renomear, Lixeira, mudanças do Claude e a trava que não deixa sair da pasta do trabalho.

As rotas rodam no manipulador de verdade do servidor.py (com o cookie de sessão), num servidor dentro do processo do
teste: assim dá para trocar os.startfile e o Popen e conferir abrir e revelar sem abrir janela nenhuma na tela.
"""
import base64
import codecs
import ctypes
import datetime as dt
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import types
import unittest
import urllib.error
import urllib.parse
import urllib.request
import wave
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest import mock

import base  # noqa: F401  (pasta de dados temporária antes de carregar o servidor)
import banco
import config
import execucao
import servidor  # o manipulador de verdade: Host, sessão, despacho das rotas
import explorador  # registra as rotas /api/arquivos (o servidor.py também importa depois de integrado)
from ferramentas import code as fcode, word


class _Http(ThreadingHTTPServer):
    daemon_threads = True

    def handle_error(self, request, client_address):
        # o Edge do teste fecha conexões no meio (ao recarregar e ao sair): não é erro do servidor
        if isinstance(sys.exc_info()[1], (ConnectionResetError, ConnectionAbortedError, BrokenPipeError)):
            return
        super().handle_error(request, client_address)


class _Srv:
    def __init__(self):
        self.http = _Http(("127.0.0.1", 0), servidor.Manipulador)
        self.http.daemon_threads = True
        self.porta = self.http.server_address[1]
        threading.Thread(target=self.http.serve_forever, kwargs={"poll_interval": 0.2}, daemon=True).start()

    def fechar(self):
        self.http.shutdown()
        self.http.server_close()

    def bruto(self, metodo, caminho, corpo=None, cab=None):
        dados = json.dumps(corpo).encode() if corpo is not None else None
        req = urllib.request.Request(f"http://127.0.0.1:{self.porta}{caminho}", data=dados, method=metodo,
                                     headers={"Content-Type": "application/json",
                                              "Cookie": f"{servidor.COOKIE}={servidor.SEGREDO}", **(cab or {})})
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.status, dict(r.headers), r.read()
        except urllib.error.HTTPError as e:
            return e.code, dict(e.headers), e.read()

    def get(self, rota, **q):
        status, _, corpo = self.bruto("GET", f"/api/arquivos/{rota}?" + urllib.parse.urlencode(q))
        return status, json.loads(corpo or b"{}")

    def post(self, rota, corpo):
        status, _, resp = self.bruto("POST", f"/api/arquivos/{rota}", corpo)
        return status, json.loads(resp or b"{}")


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = _Srv()

    @classmethod
    def tearDownClass(cls):
        cls.srv.fechar()

    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="clave-arq-")).resolve()
        self.raiz = str(self.pasta)
        self.addCleanup(shutil.rmtree, self.pasta, True)

    def get(self, rota, **q):
        return self.srv.get(rota, raiz=self.raiz, **q)

    def post(self, rota, **corpo):
        return self.srv.post(rota, {"raiz": self.raiz, **corpo})

    def escrever(self, rel, dados):
        p = self.pasta / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(dados, str):
            p.write_text(dados, encoding="utf-8", newline="")
        else:
            p.write_bytes(dados)
        return p


class TestLista(Base):
    def test_pastas_primeiro_ordem_natural_e_ocultas(self):
        for nome in ("b10.txt", "b2.txt", "A.txt", ".gitignore"):
            self.escrever(nome, "x")
        for nome in ("sub", "Docs", ".git", "node_modules", "__pycache__", ".venv", "dist"):
            (self.pasta / nome).mkdir()
        status, r = self.get("lista")
        self.assertEqual(status, 200, r)
        self.assertEqual([i["nome"] for i in r["itens"]], ["Docs", "sub", ".gitignore", "A.txt", "b2.txt", "b10.txt"])
        self.assertEqual(r["escondidos"], 5)
        self.assertEqual(r["caminho"], "")
        self.assertFalse(r["cortado"])
        arq = next(i for i in r["itens"] if i["nome"] == "A.txt")
        self.assertEqual((arq["tipo"], arq["tamanho"], arq["caminho"]), ("arquivo", 1, "A.txt"))
        self.assertGreater(arq["modificado"], 0)
        _, r = self.get("lista", ocultos=1)
        nomes = [i["nome"] for i in r["itens"]]
        for nome in (".git", "node_modules", "__pycache__", ".venv", "dist"):
            self.assertIn(nome, nomes)
        self.assertTrue(next(i for i in r["itens"] if i["nome"] == ".git")["oculto"])

    def test_subpasta_com_barra_invertida(self):
        self.escrever("sub/dentro/x.py", "print(1)")
        status, r = self.get("lista", caminho="sub\\dentro")
        self.assertEqual(status, 200, r)
        self.assertEqual(r["caminho"], "sub/dentro")
        self.assertEqual(r["itens"][0]["caminho"], "sub/dentro/x.py")

    def test_corta_em_2000(self):
        d = self.pasta / "muitos"
        d.mkdir()
        for i in range(2010):
            (d / f"a{i}.txt").write_bytes(b"")
        status, r = self.get("lista", caminho="muitos")
        self.assertEqual(status, 200)
        self.assertTrue(r["cortado"])
        self.assertEqual(len(r["itens"]), 2000)
        self.assertEqual(r["total"], 2010)
        self.assertEqual([i["nome"] for i in r["itens"][:3]], ["a0.txt", "a1.txt", "a2.txt"])

    def test_arquivo_protegido_marcado(self):
        self.escrever(".env", "SEGREDO=1")
        self.escrever("chave.pem", "x")
        _, r = self.get("lista")
        marcas = {i["nome"]: i.get("protegido") for i in r["itens"]}
        self.assertEqual(marcas[".env"], ".env")
        self.assertEqual(marcas["chave.pem"], "*.pem")

    def test_pasta_que_nao_existe(self):
        status, r = self.get("lista", caminho="nao-tem")
        self.assertEqual(status, 404)
        self.assertIn("nao-tem", r["erro"])


class TestNaoSaiDaRaiz(Base):
    def test_pontos_absolutos_e_unidade(self):
        self.escrever("sub/a.txt", "a")
        for caminho in ("..", "../", "sub/../../x", "..\\..\\Windows"):
            status, r = self.get("lista", caminho=caminho)
            self.assertEqual(status, 403, caminho)
            self.assertIn("erro", r)
        for caminho in ("C:\\Windows", "c:/windows/win.ini", "/etc/passwd", "\\\\servidor\\x", "a.txt:fluxo"):
            status, r = self.get("conteudo", caminho=caminho)
            self.assertEqual(status, 400, caminho)
        status, _ = self.get("conteudo", caminho="sub/a.txt")
        self.assertEqual(status, 200)

    def test_raiz_invalida(self):
        status, r = self.srv.get("lista", raiz="pasta-relativa")
        self.assertEqual(status, 400)
        status, r = self.srv.get("lista", raiz=str(self.pasta / "sumiu"))
        self.assertEqual(status, 404)
        status, r = self.srv.get("lista")
        self.assertEqual(status, 400)

    @unittest.skipUnless(sys.platform == "win32", "junção é coisa do Windows")
    def test_juncao_para_fora(self):
        fora = Path(tempfile.mkdtemp(prefix="clave-fora-")).resolve()
        self.addCleanup(shutil.rmtree, fora, True)
        (fora / "segredo.txt").write_text("não pode", encoding="utf-8")
        (self.pasta / "dentro").mkdir()
        (self.pasta / "dentro" / "ok.txt").write_text("ok", encoding="utf-8")
        for nome, alvo in (("link-fora", fora), ("link-dentro", self.pasta / "dentro")):
            r = subprocess.run(["cmd", "/c", "mklink", "/J", str(self.pasta / nome), str(alvo)], capture_output=True)
            if r.returncode != 0:
                self.skipTest("não deu para criar a junção")
        self.addCleanup(lambda: [os.rmdir(self.pasta / n) for n in ("link-fora", "link-dentro") if (self.pasta / n).exists()])
        _, r = self.get("lista")
        itens = {i["nome"]: i for i in r["itens"]}
        self.assertTrue(itens["link-fora"]["link"])
        self.assertTrue(itens["link-fora"].get("fora"))
        self.assertFalse(itens["link-dentro"].get("fora", False))
        self.assertEqual(self.get("lista", caminho="link-fora")[0], 403)
        self.assertEqual(self.get("conteudo", caminho="link-fora/segredo.txt")[0], 403)
        self.assertEqual(self.get("bruto", caminho="link-fora/segredo.txt")[0], 403)
        self.assertEqual(self.post("criar", pasta="link-fora", nome="novo.txt")[0], 403)
        self.assertFalse((fora / "novo.txt").exists())
        status, r = self.get("conteudo", caminho="link-dentro/ok.txt")
        self.assertEqual((status, r.get("texto")), (200, "ok"))
        # a busca não entra em link nenhum
        _, r = self.get("busca", termo="segredo")
        self.assertEqual(r["resultados"], [])
        # mandar a junção para a Lixeira levaria o que está do outro lado: recusado
        status, r = self.post("lixeira", caminho="link-dentro")
        self.assertEqual(status, 400)
        self.assertTrue((self.pasta / "dentro" / "ok.txt").exists())


class TestConteudo(Base):
    def ler(self, nome, dados):
        self.escrever(nome, dados)
        status, r = self.get("conteudo", caminho=nome)
        self.assertEqual(status, 200, r)
        return r

    def test_codificacoes(self):
        texto = "ação, coração e pão\nsegunda linha"
        casos = [("utf8.txt", texto.encode("utf-8"), "utf-8"),
                 ("bom.txt", codecs.BOM_UTF8 + texto.encode("utf-8"), "utf-8-sig"),
                 ("le.txt", codecs.BOM_UTF16_LE + texto.encode("utf-16-le"), "utf-16"),
                 ("be.txt", codecs.BOM_UTF16_BE + texto.encode("utf-16-be"), "utf-16"),
                 ("sem-bom-16.txt", texto.encode("utf-16-le"), "utf-16"),
                 ("antigo.txt", texto.encode("cp1252"), "cp1252")]
        for nome, dados, cod in casos:
            r = self.ler(nome, dados)
            self.assertFalse(r["binario"], nome)
            self.assertEqual(r["texto"], texto, nome)
            self.assertEqual(r["codificacao"], cod, nome)
            self.assertFalse(r["cortado"])

    def test_binario(self):
        r = self.ler("foto.png", b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + bytes(range(256)) * 10)
        self.assertTrue(r["binario"])
        self.assertNotIn("texto", r)
        self.assertEqual(r["mime"], "image/png")
        r = self.ler("controles.bin", bytes([1, 2, 3, 4, 5, 6, 7, 14, 15, 16]) * 100)
        self.assertTrue(r["binario"])

    def test_corta_em_2_mb_sem_quebrar_caractere(self):
        linha = ("é" * 99 + "\n").encode("utf-8")  # 199 bytes: o corte de 2 MB cai no meio de um "é"
        dados = linha * (3 * 1024 * 1024 // len(linha))
        r = self.ler("grande.txt", dados)
        self.assertTrue(r["cortado"])
        self.assertEqual(r["codificacao"], "utf-8")
        self.assertEqual(r["tamanho"], len(dados))
        self.assertNotIn("\ufffd", r["texto"])
        self.assertLessEqual(len(r["texto"].encode("utf-8")), explorador.LIMITE_TEXTO)
        self.assertGreater(len(r["texto"].encode("utf-8")), explorador.LIMITE_TEXTO - 4)

    def test_linguagem_pela_extensao(self):
        for nome, lingua in (("a.py", "python"), ("b.md", "markdown"), ("c.tsx", "typescript"), ("Makefile", "makefile"),
                             ("d.html", "xml"), ("e.json", "json"), ("f.desconhecida", "plaintext"), ("g.ps1", "powershell")):
            self.assertEqual(self.ler(nome, "x")["linguagem"], lingua, nome)

    def test_pasta_nao_tem_conteudo(self):
        (self.pasta / "sub").mkdir()
        self.assertEqual(self.get("conteudo", caminho="sub")[0], 400)

    def test_protegido_precisa_de_forcar(self):
        self.escrever(".env", "CHAVE=123")
        status, r = self.get("conteudo", caminho=".env")
        self.assertEqual(status, 403)
        self.assertEqual(r["codigo"], "protegido")
        self.assertEqual(r["padrao"], ".env")
        self.assertEqual(self.get("bruto", caminho=".env")[0], 403)
        status, r = self.get("conteudo", caminho=".env", forcar=1)
        self.assertEqual((status, r["texto"]), (200, "CHAVE=123"))
        _, info = self.get("info", caminho=".env")
        self.assertEqual(info["protegido"], ".env")


class TestBruto(Base):
    def test_range_e_cabecalhos(self):
        dados = bytes(range(256)) * 4
        self.escrever("som.mp3", dados)
        q = "?" + urllib.parse.urlencode({"raiz": self.raiz, "caminho": "som.mp3"})
        status, cab, corpo = self.srv.bruto("GET", "/api/arquivos/bruto" + q)
        self.assertEqual((status, corpo), (200, dados))
        self.assertEqual(cab["Content-Type"], "audio/mpeg")
        self.assertEqual(cab["Accept-Ranges"], "bytes")
        self.assertEqual(cab["X-Content-Type-Options"], "nosniff")
        self.assertTrue(cab["Content-Security-Policy"].startswith("sandbox"))
        self.assertTrue(cab["Content-Disposition"].startswith("inline"))
        self.assertNotIn("X-Frame-Options", cab)
        status, cab, corpo = self.srv.bruto("GET", "/api/arquivos/bruto" + q, cab={"Range": "bytes=10-19"})
        self.assertEqual((status, corpo), (206, dados[10:20]))
        self.assertEqual(cab["Content-Range"], f"bytes 10-19/{len(dados)}")
        status, _, corpo = self.srv.bruto("GET", "/api/arquivos/bruto" + q, cab={"Range": "bytes=1000-"})
        self.assertEqual((status, corpo), (206, dados[1000:]))
        status, _, corpo = self.srv.bruto("GET", "/api/arquivos/bruto" + q, cab={"Range": "bytes=-5"})
        self.assertEqual((status, corpo), (206, dados[-5:]))
        status, cab, _ = self.srv.bruto("GET", "/api/arquivos/bruto" + q, cab={"Range": "bytes=5000-"})
        self.assertEqual(status, 416)
        self.assertEqual(cab["Content-Range"], f"bytes */{len(dados)}")
        status, cab, corpo = self.srv.bruto("HEAD", "/api/arquivos/bruto" + q)
        self.assertEqual((status, corpo, cab["Content-Length"]), (200, b"", str(len(dados))))
        # o nome no fim do endereço (para o título do leitor de PDF) não muda o arquivo entregue
        status, _, corpo = self.srv.bruto("GET", "/api/arquivos/bruto/qualquer%20nome.mp3" + q, cab={"Range": "bytes=0-3"})
        self.assertEqual((status, corpo), (206, dados[:4]))

    def test_nunca_como_pagina(self):
        for nome, esperado in (("p.html", "text/plain"), ("p.htm", "text/plain"), ("x.xml", "text/plain"),
                               ("s.js", "text/plain"), ("d.svg", "image/svg+xml"), ("doc.pdf", "application/pdf"),
                               ("f.jpg", "image/jpeg"), ("v.mp4", "video/mp4"), ("prog.exe", "application/octet-stream")):
            self.escrever(nome, b"<html><script>alert(1)</script></html>")
            q = "?" + urllib.parse.urlencode({"raiz": self.raiz, "caminho": nome})
            status, cab, _ = self.srv.bruto("GET", "/api/arquivos/bruto" + q)
            self.assertEqual(status, 200)
            self.assertEqual(cab["Content-Type"], esperado, nome)
            self.assertNotIn("html", cab["Content-Type"])
            self.assertTrue(cab["Content-Security-Policy"].startswith("sandbox"), nome)
            self.assertIn("default-src 'none'", cab["Content-Security-Policy"])


class TestPlanilhas(Base):
    def test_xlsx_abas_formula_data_e_corte(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vendas"
        ws.append(["Produto", "Preço", "Quando"])
        ws.append(["Café", 12.5, dt.datetime(2026, 10, 15)])
        ws.append(["Pão", 3, dt.datetime(2026, 10, 15, 8, 30)])
        ws["B4"] = "=SUM(B2:B3)"
        ws["D1"] = True
        ws2 = wb.create_sheet("Longa")
        for i in range(250):
            ws2.append([i] + [f"c{j}" for j in range(60)])
        wb.save(self.pasta / "planilha.xlsx")
        status, r = self.get("planilha", caminho="planilha.xlsx")
        self.assertEqual(status, 200, r)
        self.assertEqual(r["tipo"], "xlsx")
        self.assertEqual([a["nome"] for a in r["abas"]], ["Vendas", "Longa"])
        v = r["abas"][0]
        self.assertEqual(v["linhas"][0], ["Produto", "Preço", "Quando", "VERDADEIRO"])
        self.assertEqual(v["linhas"][1], ["Café", "12.5", "2026-10-15"])
        self.assertEqual(v["linhas"][2], ["Pão", "3", "2026-10-15 08:30"])
        self.assertEqual(v["linhas"][3], ["", "=SUM(B2:B3)"], "fórmula nunca calculada aparece como fórmula")
        self.assertFalse(v["cortado"])
        longa = r["abas"][1]
        self.assertTrue(longa["cortado"])
        self.assertEqual(len(longa["linhas"]), 200)
        self.assertEqual(len(longa["linhas"][0]), 50)
        self.assertEqual(longa["total_linhas"], 250)

    def test_csv_com_ponto_e_virgula_em_cp1252(self):
        self.escrever("dados.csv", "nome;cidade;valor\nJoão;São Paulo;1,50\nAna;\"Belém; PA\";2\n".encode("cp1252"))
        status, r = self.get("planilha", caminho="dados.csv")
        self.assertEqual(status, 200, r)
        aba = r["abas"][0]
        self.assertEqual(aba["separador"], ";")
        self.assertEqual(aba["linhas"], [["nome", "cidade", "valor"], ["João", "São Paulo", "1,50"], ["Ana", "Belém; PA", "2"]])

    def test_csv_com_virgula_e_tsv(self):
        self.escrever("a.csv", "a,b,c\n1,2,3\n")
        self.assertEqual(self.get("planilha", caminho="a.csv")[1]["abas"][0]["linhas"], [["a", "b", "c"], ["1", "2", "3"]])
        self.escrever("b.tsv", "x\ty\n1\t2\n")
        self.assertEqual(self.get("planilha", caminho="b.tsv")[1]["abas"][0]["linhas"], [["x", "y"], ["1", "2"]])

    def test_nao_e_planilha(self):
        self.escrever("velha.xls", b"\xd0\xcf\x11\xe0")
        status, r = self.get("planilha", caminho="velha.xls")
        self.assertEqual(status, 400)
        self.assertIn(".xls", r["erro"])
        self.escrever("quebrada.xlsx", b"nao sou zip")
        status, r = self.get("planilha", caminho="quebrada.xlsx")
        self.assertEqual(status, 400)


class TestDocx(Base):
    def test_texto_do_documento(self):
        word.criar(self.pasta / "carta.docx", "# Título\n\nOlá, **mundo**.\n\n| A | B |\n|---|---|\n| 1 | 2 |")
        status, r = self.get("docx", caminho="carta.docx")
        self.assertEqual(status, 200, r)
        self.assertIn("# Título", r["texto"])
        self.assertIn("Olá, mundo.", r["texto"])
        self.assertIn("| 1 | 2 |", r["texto"])
        self.escrever("falso.docx", b"nao sou docx")
        status, r = self.get("docx", caminho="falso.docx")
        self.assertEqual(status, 400)
        self.assertIn(".docx", r["erro"])


class TestMudar(Base):
    def test_criar(self):
        status, r = self.post("criar", pasta="", nome="novo.txt", tipo="arquivo")
        self.assertEqual((status, r["caminho"]), (200, "novo.txt"))
        self.assertTrue((self.pasta / "novo.txt").is_file())
        status, r = self.post("criar", pasta="", nome="src", tipo="pasta")
        self.assertEqual((status, r["caminho"]), (200, "src"))
        status, r = self.post("criar", pasta="src", nome="app.js", tipo="arquivo")
        self.assertEqual((status, r["caminho"]), (200, "src/app.js"))
        self.assertEqual(self.post("criar", pasta="", nome="novo.txt")[0], 409)
        for ruim in ("CON", "nul.txt", "a:b", "a/b", "fim.", "", "..", "x?"):
            status, r = self.post("criar", pasta="", nome=ruim)
            self.assertEqual(status, 400, ruim)
            self.assertIn("erro", r)

    def test_renomear(self):
        self.escrever("velho.txt", "x")
        self.escrever("outro.txt", "y")
        status, r = self.post("renomear", caminho="velho.txt", nome="novo.txt")
        self.assertEqual((status, r["caminho"]), (200, "novo.txt"))
        self.assertTrue((self.pasta / "novo.txt").is_file())
        self.assertFalse((self.pasta / "velho.txt").exists())
        self.assertEqual(self.post("renomear", caminho="novo.txt", nome="outro.txt")[0], 409)
        status, r = self.post("renomear", caminho="novo.txt", nome="Novo.txt")  # só a caixa das letras
        self.assertEqual(status, 200, r)
        self.assertIn("Novo.txt", os.listdir(self.pasta))
        self.assertEqual(self.post("renomear", caminho="", nome="x")[0], 400)
        self.assertEqual(self.post("renomear", caminho="nao-tem.txt", nome="x.txt")[0], 404)

    @unittest.skipUnless(sys.platform == "win32", "Lixeira do Windows")
    def test_lixeira(self):
        p = self.escrever("clave-teste-lixeira.txt", "pode ir para a Lixeira")
        status, r = self.post("lixeira", caminho="clave-teste-lixeira.txt")
        self.assertEqual(status, 200, r)
        self.assertFalse(p.exists())
        (self.pasta / "clave-teste-pasta-lixeira").mkdir()
        self.escrever("clave-teste-pasta-lixeira/dentro.txt", "x")
        status, r = self.post("lixeira", caminho="clave-teste-pasta-lixeira")
        self.assertEqual(status, 200, r)
        self.assertFalse((self.pasta / "clave-teste-pasta-lixeira").exists())
        self.assertEqual(self.post("lixeira", caminho="")[0], 400)
        self.assertTrue(self.pasta.is_dir())
        self.assertEqual(self.post("lixeira", caminho="../fora.txt")[0], 403)

    def test_abrir_e_revelar_sem_abrir_janela(self):
        self.escrever("leia.txt", "x")
        self.escrever("roda.bat", "echo oi")
        with mock.patch.object(explorador.os, "startfile", create=True) as abrir, \
                mock.patch.object(explorador.subprocess, "Popen") as popen:
            status, r = self.post("abrir", caminho="leia.txt")
            self.assertEqual(status, 200, r)
            abrir.assert_called_once_with(str(self.pasta / "leia.txt"))
            status, r = self.post("abrir", caminho="roda.bat")
            self.assertEqual((status, r["codigo"]), (409, "executavel"))
            self.assertEqual(abrir.call_count, 1, "programa só abre com confirmação")
            status, r = self.post("abrir", caminho="roda.bat", confirmado=True)
            self.assertEqual(status, 200)
            self.assertEqual(abrir.call_count, 2)
            status, r = self.post("revelar", caminho="leia.txt")
            self.assertEqual(status, 200, r)
            if sys.platform == "win32":
                self.assertEqual(popen.call_args[0][0], f'explorer.exe /select,"{self.pasta / "leia.txt"}"')
            self.assertEqual(self.post("abrir", caminho="..\\..\\Windows\\notepad.exe", confirmado=True)[0], 403)
            self.assertEqual(abrir.call_count, 2)


class TestBuscaEInfo(Base):
    def test_busca_pelo_nome(self):
        self.escrever("a/b/c/alvo-fundo.py", "x")
        self.escrever("alvo.txt", "x")
        self.escrever("node_modules/pacote/alvo.js", "x")
        self.escrever("outro.md", "x")
        status, r = self.get("busca", termo="ALVO")
        self.assertEqual(status, 200)
        caminhos = [x["caminho"] for x in r["resultados"]]
        self.assertEqual(caminhos, ["alvo.txt", "a/b/c/alvo-fundo.py"], "mais rasos primeiro, sem node_modules")
        _, r = self.get("busca", termo="*.py")
        self.assertEqual([x["caminho"] for x in r["resultados"]], ["a/b/c/alvo-fundo.py"])
        _, r = self.get("busca", termo="b")
        self.assertIn({"nome": "b", "caminho": "a/b", "tipo": "pasta"}, r["resultados"])
        self.assertEqual(self.get("busca", termo="")[0], 400)

    def test_busca_para_em_200(self):
        for i in range(230):
            (self.pasta / f"item{i}.txt").write_bytes(b"")
        _, r = self.get("busca", termo="item")
        self.assertEqual(len(r["resultados"]), 200)
        self.assertTrue(r["cortado"])

    def test_info_e_mtimes(self):
        p = self.escrever("x.py", "print(1)")
        status, r = self.get("info", caminho="x.py")
        self.assertEqual(status, 200)
        self.assertEqual((r["tipo"], r["tamanho"], r["linguagem"], r["absoluto"]), ("arquivo", 8, "python", str(p)))
        self.assertTrue(r["executavel"], "um .py abre no Python: roda")
        self.assertIsNone(r["protegido"])
        self.escrever("nota.txt", "x")
        self.assertFalse(self.get("info", caminho="nota.txt")[1]["executavel"])
        _, r = self.get("info", caminho="")
        self.assertEqual((r["tipo"], r["caminho"]), ("pasta", ""))
        os.utime(p, (1_700_000_000, 1_700_000_000))
        status, r = self.post("mtimes", caminhos=["x.py", "", "nao-tem.txt", "../fora"])
        self.assertEqual(status, 200)
        self.assertEqual(r["mtimes"]["x.py"], 1_700_000_000)
        self.assertIsNotNone(r["mtimes"][""])
        self.assertIsNone(r["mtimes"]["nao-tem.txt"])
        self.assertIsNone(r["mtimes"]["../fora"])


class TestMudancasDoClaude(Base):
    def test_antes_e_depois_com_checkpoint_de_verdade(self):
        p = self.escrever("app.py", "a = 1\nb = 2\n")
        conv = banco.criar_conversa("code", pasta=self.raiz)
        eid = banco.novo_id()
        ctx = fcode.Contexto(raiz=self.pasta, protegidos=[], execucao_id=eid, conversa_id=conv["id"])
        fcode.executar("ler_arquivo", {"caminho": "app.py"}, ctx)
        fcode.executar("editar_arquivo", {"caminho": "app.py", "trecho_atual": "b = 2", "trecho_novo": "b = 3"}, ctx)
        fcode.executar("escrever_arquivo", {"caminho": "novo/criado.txt", "conteudo": "oi\n"}, ctx)
        self.assertEqual(p.read_text(encoding="utf-8"), "a = 1\nb = 3\n")

        status, r = self.get("mudancas", conversa=conv["id"])
        self.assertEqual(status, 200, r)
        self.assertEqual(sorted(r["arquivos"]), ["app.py", "novo/criado.txt"])

        status, r = self.get("mudancas", conversa=conv["id"], caminho="app.py")
        self.assertEqual(status, 200, r)
        self.assertTrue(r["mudou"])
        self.assertFalse(r["criado"] or r["igual"] or r["binario"])
        self.assertEqual((r["antes"], r["depois"]), ("a = 1\nb = 2\n", "a = 1\nb = 3\n"))
        self.assertIn("-b = 2", r["diff"])
        self.assertIn("+b = 3", r["diff"])

        _, r = self.get("mudancas", conversa=conv["id"], caminho="novo/criado.txt")
        self.assertTrue(r["criado"])
        self.assertEqual((r["antes"], r["depois"]), ("", "oi\n"))

        _, info = self.get("info", caminho="app.py", conversa=conv["id"])
        self.assertTrue(info["mudou"])
        self.escrever("intocado.txt", "x")
        _, info = self.get("info", caminho="intocado.txt", conversa=conv["id"])
        self.assertFalse(info["mudou"])
        _, r = self.get("mudancas", conversa=conv["id"], caminho="intocado.txt")
        self.assertFalse(r["mudou"])

        # outra conversa não vê as mudanças desta
        outra = banco.criar_conversa("code", pasta=self.raiz)
        self.assertEqual(self.get("mudancas", conversa=outra["id"])[1]["arquivos"], [])

        execucao.desfazer(eid)
        _, r = self.get("mudancas", conversa=conv["id"], caminho="app.py")
        self.assertTrue(r["igual"], "depois de desfazer, agora é igual a antes")
        self.assertEqual(r["diff"], "")
        _, r = self.get("mudancas", conversa=conv["id"], caminho="novo/criado.txt")
        self.assertTrue(r["apagado"])
        self.assertEqual(self.get("mudancas", caminho="app.py")[0], 400)


def _checkpoint(conversa_id: str, p: Path, conteudo: bytes | None, existia: bool = True) -> str:
    eid = banco.novo_id()
    banco.guardar_checkpoint(eid, conversa_id, str(p), existia, conteudo)
    return eid


class TestMudancasSemLerDemais(Base):
    def test_lista_e_info_nao_leem_o_conteudo_de_todos_os_checkpoints(self):
        conv = banco.criar_conversa("code", pasta=self.raiz)
        for i in range(3):
            p = self.escrever(f"a{i}.txt", f"agora {i}\n")
            _checkpoint(conv["id"], p, f"antes {i}\n".encode() + b"x" * 200_000)
        sqls: list[str] = []
        original = banco.conexao

        def conexao():
            con = original()
            con.set_trace_callback(sqls.append)
            return con

        with mock.patch.object(banco, "conexao", conexao):
            status, r = self.get("mudancas", conversa=conv["id"])
            self.assertEqual((status, sorted(r["arquivos"])), (200, ["a0.txt", "a1.txt", "a2.txt"]))
            status, r = self.get("info", caminho="a1.txt", conversa=conv["id"])
            self.assertEqual((status, r["mudou"]), (200, True))
            leram_lista = [s for s in sqls if "checkpoints" in s and "conteudo" in s]
            self.assertEqual(leram_lista, [], "a lista e a info não precisam do conteúdo de nenhum checkpoint")
            status, r = self.get("mudancas", conversa=conv["id"], caminho="a1.txt")
        self.assertEqual(status, 200, r)
        self.assertTrue(r["antes"].startswith("antes 1\n"))
        self.assertEqual(r["depois"], "agora 1\n")
        com_conteudo = [s for s in sqls if "checkpoints" in s and "conteudo" in s and "length(" not in s]
        self.assertEqual(len(com_conteudo), 1, com_conteudo)
        self.assertIn("WHERE id=", com_conteudo[0], "só o conteúdo do arquivo pedido")

    def test_arquivo_grande_igual_tem_tamanho_de_verdade(self):
        dados = bytes(range(256)) * (3 * 1024 * 1024 // 256)  # 3 MB, maior que o limite da tela (2 MB)
        p = self.escrever("grande.bin", dados)
        conv = banco.criar_conversa("code", pasta=self.raiz)
        _checkpoint(conv["id"], p, dados)
        status, r = self.get("mudancas", conversa=conv["id"], caminho="grande.bin")
        self.assertEqual(status, 200, r)
        self.assertTrue(r["grande"])
        self.assertTrue(r["igual"], "igual ao de antes: não pode parecer mudado só porque é grande")
        self.assertEqual((r["tamanho_antes"], r["tamanho_depois"]), (len(dados), len(dados)))
        # mesmo tamanho, último byte diferente
        p.write_bytes(dados[:-1] + b"!")
        _, r = self.get("mudancas", conversa=conv["id"], caminho="grande.bin")
        self.assertFalse(r["igual"])
        self.assertEqual(r["tamanho_depois"], len(dados))
        # cresceu: o tamanho é o de agora, não o do corte
        p.write_bytes(dados + b"mais")
        _, r = self.get("mudancas", conversa=conv["id"], caminho="grande.bin")
        self.assertFalse(r["igual"])
        self.assertEqual(r["tamanho_depois"], len(dados) + 4)
        # apagado
        p.unlink()
        _, r = self.get("mudancas", conversa=conv["id"], caminho="grande.bin")
        self.assertEqual((r["igual"], r["apagado"], r["tamanho_depois"]), (False, True, 0))

    def test_antes_grande_e_agora_pequeno(self):
        p = self.escrever("encolheu.txt", "pouco\n")
        conv = banco.criar_conversa("code", pasta=self.raiz)
        _checkpoint(conv["id"], p, b"a\n" * (1024 * 1024 + 10))
        _, r = self.get("mudancas", conversa=conv["id"], caminho="encolheu.txt")
        self.assertEqual((r["grande"], r["igual"], r["tamanho_depois"]), (True, False, 6))
        self.assertEqual(r["tamanho_antes"], 2 * (1024 * 1024 + 10))

    def test_pedacos_do_checkpoint(self):
        conv = banco.criar_conversa("code", pasta=self.raiz)
        p = self.pasta / "x.bin"
        dados = os.urandom(300_000)
        _checkpoint(conv["id"], p, dados)
        ck = explorador._primeiro_checkpoint(conv["id"], p)
        self.assertNotIn("conteudo", ck)
        self.assertEqual(explorador._tamanho_checkpoint(ck["id"]), len(dados))
        self.assertEqual(b"".join(explorador._pedacos_checkpoint(ck["id"], len(dados), bloco=70_000)), dados)
        self.assertEqual(explorador._conteudo_checkpoint(ck["id"]), dados)
        self.assertTrue(explorador._mesmo_conteudo(io.BytesIO(dados), explorador._pedacos_checkpoint(ck["id"], len(dados))))
        self.assertFalse(explorador._mesmo_conteudo(io.BytesIO(dados + b"!"), explorador._pedacos_checkpoint(ck["id"], len(dados))))


class TestExecutaveis(Base):
    def test_extensao_como_o_windows_ve(self):
        e = explorador._extensao
        self.assertEqual([e("roda.bat."), e("roda.bat "), e("roda.BAT. . "), e(".bat"), e("a.tar.gz"), e("sem"), e("")],
                         [".bat", ".bat", ".bat", ".bat", ".gz", "", ""])

    def test_formatos_de_macro_e_programa(self):
        for ext in (".pyc", ".pyo", ".pyw", ".psc1", ".diagcab", ".msu", ".msp", ".appref-ms", ".xll", ".xlsb", ".xla",
                    ".xlam", ".dotm", ".ppsm", ".potm", ".docm", ".xlsm", ".pptm", ".sldm", ".PS1", ".accde", ".wsb"):
            self.escrever("arquivo" + ext, "x")
            status, info = self.get("info", caminho="arquivo" + ext)
            self.assertEqual(status, 200, info)
            self.assertTrue(info["executavel"], ext)
        for nome in ("nota.txt", "foto.png", "planilha.xlsx", "texto.docx", "dados.csv", "pagina.html", "video.mp4"):
            self.escrever(nome, "x")
            self.assertFalse(self.get("info", caminho=nome)[1]["executavel"], nome)

    @unittest.skipUnless(sys.platform == "win32", "o Windows tira ponto e espaço do fim do nome")
    def test_ponto_ou_espaco_no_fim_nao_pula_a_confirmacao(self):
        self.escrever("roda.bat", "echo oi")
        with mock.patch.object(explorador.os, "startfile", create=True) as abrir:
            for nome in ("roda.bat.", "roda.bat ", "roda.bat. .", "RODA.BAT"):
                status, r = self.post("abrir", caminho=nome)
                self.assertEqual((status, r.get("codigo")), (409, "executavel"), nome)
                self.assertTrue(self.get("info", caminho=nome)[1]["executavel"], nome)
            self.assertEqual(abrir.call_count, 0, "nada abriu sem a confirmação")
            status, _ = self.post("abrir", caminho="roda.bat.", confirmado=True)
            self.assertEqual((status, abrir.call_count), (200, 1))

    @unittest.skipUnless(sys.platform == "win32", "nome curto 8.3 é coisa do Windows")
    def test_nome_curto_nao_pula_a_confirmacao(self):
        p = self.escrever("programa-comprido.appref-ms", "x")
        buf = ctypes.create_unicode_buffer(1024)
        if not ctypes.windll.kernel32.GetShortPathNameW(str(p), buf, 1024) or Path(buf.value).name == p.name:
            self.skipTest("este disco não guarda nomes curtos")
        curto = Path(buf.value).name  # PROGRA~1.APP: a extensão curta não está na lista
        self.assertNotIn(explorador._extensao(curto), explorador.EXECUTAVEIS)
        with mock.patch.object(explorador.os, "startfile", create=True) as abrir:
            status, r = self.post("abrir", caminho=curto)
            self.assertEqual((status, r.get("codigo")), (409, "executavel"))
            self.assertEqual(abrir.call_count, 0)


@unittest.skipUnless(sys.platform == "win32", "atributos oculto e de sistema do Windows")
class TestBuscaEscondeOQueAArvoreEsconde(Base):
    def _atributos(self, p: Path, valor: int):
        self.assertTrue(ctypes.windll.kernel32.SetFileAttributesW(str(p), valor))

    def test_oculto_e_de_sistema(self):
        self.escrever("alvo-visivel.txt", "x")
        so_oculto = self.escrever("alvo-so-oculto.txt", "x")
        sistema = self.escrever("alvo-sistema.txt", "x")
        pasta = self.pasta / "pasta-sistema"
        pasta.mkdir()
        self.escrever("pasta-sistema/alvo-dentro.txt", "x")
        OCULTO, SISTEMA, NORMAL = 0x2, 0x4, 0x80
        self._atributos(so_oculto, OCULTO)
        self._atributos(sistema, OCULTO | SISTEMA)
        self._atributos(pasta, OCULTO | SISTEMA)
        self.addCleanup(lambda: [ctypes.windll.kernel32.SetFileAttributesW(str(x), NORMAL) for x in (so_oculto, sistema, pasta)])
        _, lista = self.get("lista")
        self.assertEqual(sorted(i["nome"] for i in lista["itens"]), ["alvo-so-oculto.txt", "alvo-visivel.txt"])
        _, r = self.get("busca", termo="alvo")
        self.assertEqual(sorted(x["caminho"] for x in r["resultados"]), ["alvo-so-oculto.txt", "alvo-visivel.txt"],
                         "a busca mostra o mesmo que a árvore")
        _, r = self.get("busca", termo="alvo", ocultos=1)
        self.assertEqual(sorted(x["caminho"] for x in r["resultados"]),
                         ["alvo-sistema.txt", "alvo-so-oculto.txt", "alvo-visivel.txt", "pasta-sistema/alvo-dentro.txt"])


@unittest.skipUnless(sys.platform == "win32", "Lixeira do Windows")
class TestLixeiraComDona(unittest.TestCase):
    def test_aviso_de_apagar_de_vez_fica_preso_a_janela_da_frente(self):
        capturado = {}

        def operacao(ref):
            op = ref._obj
            capturado.update(hwnd=op.hwnd, flags=op.fFlags)
            return 0

        with mock.patch.object(explorador, "_shell32", types.SimpleNamespace(SHFileOperationW=operacao)), \
                mock.patch.object(explorador, "_user32", types.SimpleNamespace(GetForegroundWindow=lambda: 0x1234)):
            explorador.mandar_para_lixeira(Path(tempfile.gettempdir()) / "clave-nao-existe.txt")
        self.assertEqual(capturado["hwnd"], 0x1234)
        self.assertTrue(capturado["flags"] & explorador.FOF_WANTNUKEWARNING)
        # de verdade: devolve a janela da frente (ou None), sem erro
        dona = explorador._janela_dona()
        self.assertTrue(dona is None or isinstance(dona, int))


# ======================= a tela (JavaScript) =======================

def _trecho_puro() -> str:
    fonte = (base.RAIZ / "app" / "js" / "arquivos.js").read_text(encoding="utf-8")
    return fonte[fonte.index("// PURAS-INICIO"):fonte.index("// PURAS-FIM")]


class TestFuncoesPurasNoNode(unittest.TestCase):
    """As funções sem DOM do arquivos.js (o trecho entre PURAS-INICIO e PURAS-FIM), rodando no Node."""

    @classmethod
    def setUpClass(cls):
        cls.node = shutil.which("node")
        if not cls.node:
            raise unittest.SkipTest("sem o Node")

    def rodar(self, corpo: str):
        script = _trecho_puro() + "\nconst r = (() => {" + corpo + "})();\nprocess.stdout.write(JSON.stringify(r));\n"
        p = subprocess.run([self.node, "--input-type=module"], input=script, capture_output=True, text=True,
                           encoding="utf-8", timeout=60)
        self.assertEqual(p.returncode, 0, p.stderr)
        return json.loads(p.stdout)

    def test_json_organizado_mantem_numeros_ordem_e_escapes(self):
        original = '{"id":1234567890123456789,"preco":10.50,"b":1,"2":2,"t":"\\u00e9 \\"x\\" \\\\","e":1E400,"v":[],"o":{ },"n":[null,true,-0.0]}'
        r = self.rodar("const s = " + json.dumps(original) + "; JSON.parse(s); return organizarJson(s);")
        self.assertEqual(r, '{\n  "id": 1234567890123456789,\n  "preco": 10.50,\n  "b": 1,\n  "2": 2,\n'
                            '  "t": "\\u00e9 \\"x\\" \\\\",\n  "e": 1E400,\n  "v": [],\n  "o": {},\n  "n": [\n    null,\n'
                            '    true,\n    -0.0\n  ]\n}')

    def test_json_organizado_igual_ao_stringify_quando_nada_se_perde(self):
        r = self.rodar(r"""
          const casos = [{a: [1, {b: 'c', d: []}], e: {}}, [], {}, 'só texto', 12, null, [[[]]], {x: {y: {z: [1, 2, 3]}}}];
          return casos.every((v) => [JSON.stringify(v), JSON.stringify(v, null, 2), ' \r\n' + JSON.stringify(v, null, '\t') + '\n']
            .every((s) => organizarJson(s) === JSON.stringify(v, null, 2)));
        """)
        self.assertTrue(r)

    def test_caminho_com_porcentagem_no_markdown(self):
        r = self.rodar(r"""
          return [resolverRelativo('docs', 'img/100%.png'), resolverRelativo('docs', '../a%20b.md#titulo'),
                  resolverRelativo('', '100%25.md'), resolverRelativo('', '50% e%20mais.md?x=1'),
                  resolverRelativo('a/b', '/raiz.md'), decodificarCaminho('%E9 e %C3%A9')];
        """)
        self.assertEqual(r, ["docs/img/100%.png", "a b.md", "100%.md", "50% e mais.md", "raiz.md", "%E9 e é"])

    def test_numero_da_planilha_no_jeito_brasileiro_e_com_moeda(self):
        numeros = ["42", "-7", "+3", "3.14", "42,50", "R$ 42,50", "R$ 1.234,56", "-R$ 10,00", "R$ -10,00", "US$ 1,234.56",
                   "$3.50", "€ 10", "10,00 €", "15%", "12,5 %", "1.234.567,89", "1,234,567.89", "1 234,56", "1e5", ",5", "2026"]
        textos = ["", "  ", "R$", "%", "2026-10-13", "12 pães", "3 pães, 1 bolo", "=SOMA(A1:A3)", "(31) 9999-8888", "1.2.3",
                  "abc", "1,2,3", "R$ 4,5,6", "12:30"]
        r = self.rodar("return [" + ", ".join(f"pareceNumero({json.dumps(v)})" for v in numeros + textos) + "];")
        self.assertEqual(r[:len(numeros)], [True] * len(numeros), list(zip(numeros, r)))
        self.assertEqual(r[len(numeros):], [False] * len(textos), list(zip(textos, r[len(numeros):])))


def _edge() -> str | None:
    for var in ("ProgramFiles(x86)", "ProgramFiles"):
        p = Path(os.environ.get(var, "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe"
        if os.environ.get(var) and p.is_file():
            return str(p)
    return None


def _cliente_ws():
    spec = importlib.util.spec_from_file_location(
        "capturar_print", base.RAIZ / "modelo-de-pasta" / ".claude" / "referencias-visuais" / "capturar-print.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.WS


_PRELUDIO_JS = r"""
const esperar = async (f, ms = 15000) => {
  const fim = Date.now() + ms;
  for (;;) {
    let v = null;
    try { v = f(); } catch { v = null; }
    if (v) return v;
    if (Date.now() > fim) throw new Error('tempo esgotado esperando ' + f);
    await new Promise((r) => setTimeout(r, 50));
  }
};
const dormir = (ms) => new Promise((r) => setTimeout(r, ms));
const P = await import('/js/painel.js');
const A = await import('/js/arquivos.js');
const vis = () => document.querySelector('#painel-lateral .aba-visualizador');
const arq = () => document.querySelector('#painel-lateral .aba-arquivos');
const guia = (caminho) => [...vis().querySelectorAll('.vis-guia-nome')].find((b) => b.title === caminho);
const ctx = { modo: 'code', pasta: RAIZ, conversaId: null };
P.definirContexto(ctx);
await dormir(100);
if (P.contextoAtual()?.pasta !== RAIZ) P.definirContexto(ctx);
"""


@unittest.skipUnless(sys.platform == "win32" and _edge(), "precisa do Edge (Windows)")
class TestTelaNoEdge(Base):
    """A tela de verdade (app/js/arquivos.js) no Edge sem janela, falando com o servidor do teste."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.perfil = tempfile.mkdtemp(prefix="clave-teste-edge-")
        cls.proc = subprocess.Popen(
            [_edge(), "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check", "--disable-sync",
             "--mute-audio", "--disable-extensions", "--disable-features=Translate,MediaRouter",
             "--remote-debugging-port=0", f"--user-data-dir={cls.perfil}", "--window-size=1400,900", "about:blank"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            arq = Path(cls.perfil) / "DevToolsActivePort"
            limite = time.monotonic() + 40
            while not (arq.is_file() and arq.read_text(encoding="utf-8").strip()):
                if time.monotonic() > limite:
                    raise RuntimeError("o Edge não abriu a porta de depuração")
                time.sleep(0.2)
            porta = int(arq.read_text(encoding="utf-8").splitlines()[0])
            alvos = json.loads(urllib.request.urlopen(f"http://127.0.0.1:{porta}/json/list", timeout=20).read())
            pagina = next(a for a in alvos if a.get("type") == "page")
            cls.ws = _cliente_ws()(pagina["webSocketDebuggerUrl"])
        except Exception:
            cls._fechar_edge()
            raise

    @classmethod
    def _fechar_edge(cls):
        if getattr(cls, "ws", None):
            cls.ws.fechar()
        subprocess.run(["taskkill", "/T", "/F", "/PID", str(cls.proc.pid)], capture_output=True, creationflags=0x08000000)
        try:
            cls.proc.wait(timeout=15)
        except subprocess.TimeoutExpired:
            pass
        shutil.rmtree(cls.perfil, ignore_errors=True)

    @classmethod
    def tearDownClass(cls):
        cls._fechar_edge()
        super().tearDownClass()

    def avaliar(self, expressao: str, timeout=90):
        r = self.ws.cmd("Runtime.evaluate", {"expression": expressao, "awaitPromise": True, "returnByValue": True},
                        timeout=timeout)
        if r.get("exceptionDetails"):
            d = r["exceptionDetails"]
            raise AssertionError((d.get("exception") or {}).get("description") or d.get("text"))
        return (r.get("result") or {}).get("value")

    def esperar_js(self, condicao: str, segundos=30):
        limite = time.monotonic() + segundos
        while time.monotonic() < limite:
            try:
                if self.avaliar(condicao, timeout=10):
                    return
            except (AssertionError, RuntimeError, TimeoutError, OSError):
                pass
            time.sleep(0.2)
        self.fail(f"a tela não chegou em: {condicao}")

    def setUp(self):
        super().setUp()
        # app novo a cada teste, numa tela sem conversa (nenhuma outra tela muda o contexto do painel depois)
        self.ws.cmd("Page.navigate", {"url": f"http://127.0.0.1:{self.srv.porta}/?t={servidor.SEGREDO}"})
        self.esperar_js("document.readyState === 'complete' && !!document.querySelector('#painel-lateral .painel-corpo')"
                        " && !!document.querySelector('#lateral nav, #lateral button')")
        time.sleep(0.8)
        self.avaliar("location.hash = '#/uso'; 1")
        self.esperar_js("location.hash === '#/uso'")
        time.sleep(0.8)

    def na_tela(self, corpo: str, timeout=90):
        return self.avaliar("(async () => { const RAIZ = " + json.dumps(self.raiz) + ";\n" + _PRELUDIO_JS + corpo
                            + "\n})()", timeout=timeout)

    def test_json_organizado_mostra_os_numeros_como_estao(self):
        self.escrever("dados.json", '{"id":1234567890123456789,"preco":10.50,"b":1,"2":2,"nome":"\\u00e9"}')
        r = self.na_tela(r"""
          A.abrirArquivo('dados.json');
          const cod = await esperar(() => vis()?.querySelector('.vis-codigo'));
          return { modo: vis().querySelector('.vis-seg [aria-pressed="true"]')?.textContent,
                   texto: [...cod.querySelectorAll('.lin .c')].map((x) => x.textContent).join('\n') };
        """)
        self.assertEqual(r["modo"], "Organizado")
        self.assertEqual(r["texto"], '{\n  "id": 1234567890123456789,\n  "preco": 10.50,\n  "b": 1,\n  "2": 2,\n'
                                     '  "nome": "\\u00e9"\n}')

    def test_markdown_com_porcentagem_no_caminho_abre(self):
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
        self.escrever("docs/img/100%.png", png)
        self.escrever("docs/100%.md", "# Cem por cento\n")
        self.escrever("docs/notas.md", "# Notas\n\n![grafico](img/100%.png)\n\n[abre o outro](100%.md)\n")
        r = self.na_tela(r"""
          A.abrirArquivo('docs/notas.md');
          const md = await esperar(() => vis()?.querySelector('.vis-md') || vis()?.querySelector('.vis-mensagem'));
          if (!md.classList.contains('vis-md')) return { erro: md.textContent };
          const img = md.querySelector('img');
          await esperar(() => img.complete);
          const largura = img.naturalWidth, src = img.getAttribute('src');
          md.querySelector('a').click();
          await esperar(() => guia('docs/100%.md'));
          const h1 = await esperar(() => vis().querySelector('.vis-md h1')?.textContent === 'Cem por cento');
          return { largura, src, h1 };
        """)
        self.assertNotIn("erro", r, r)
        self.assertIn("caminho=docs%2Fimg%2F100%25.png", r["src"])
        self.assertEqual(r["largura"], 1, "a imagem com % no nome carregou")
        self.assertTrue(r["h1"], "o link com % abriu o outro arquivo")

    def test_comando_sem_lista_so_recarrega_o_que_mudou(self):
        buf = io.BytesIO()
        with wave.open(buf, "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(8000)
            w.writeframes(b"\0\0" * 8000)
        self.escrever("som.wav", buf.getvalue())
        nota = self.escrever("nota.txt", "linha 1\n")
        r = self.na_tela(r"""
          A.abrirArquivo('nota.txt');
          window.__cod = await esperar(() => vis()?.querySelector('.vis-codigo'));
          A.abrirArquivo('som.wav');
          const midia = await esperar(() => vis()?.querySelector('.vis-midia audio'));
          await esperar(() => midia.readyState >= 1);
          window.__midia = midia;
          dispatchEvent(new CustomEvent('clave:arquivos-mudaram', { detail: { caminhos: [] } }));
          await dormir(1500);
          const aviso = vis().querySelector('.vis-aviso');
          const mesmaMidia = vis().querySelector('.vis-midia audio') === midia;
          const avisoTexto = aviso.hidden ? '' : aviso.textContent;
          guia('nota.txt').click();
          await dormir(300);
          return { mesmaMidia, avisoTexto, mesmoCodigo: vis().querySelector('.vis-codigo') === window.__cod };
        """)
        self.assertEqual(r["avisoTexto"], "", "nada mudou: nenhum aviso de atualizado")
        self.assertTrue(r["mesmaMidia"], "o áudio que estava à vista não foi recarregado (não volta ao começo)")
        self.assertTrue(r["mesmoCodigo"], "a guia de fora da vista também não")
        # agora o comando mudou o texto (a guia do áudio à vista): só ele lê de novo
        r = self.na_tela(r"""
          guia('som.wav').click();
          await esperar(() => vis().querySelector('.vis-midia audio') === window.__midia);
          return true;
        """)
        nota.write_text("linha 2\n", encoding="utf-8")
        os.utime(nota, (time.time() + 30, time.time() + 30))
        r = self.na_tela(r"""
          dispatchEvent(new CustomEvent('clave:arquivos-mudaram', { detail: { caminhos: [] } }));
          await dormir(1500);
          const mesmaMidia = vis().querySelector('.vis-midia audio') === window.__midia;
          guia('nota.txt').click();
          const cod = await esperar(() => { const c = vis().querySelector('.vis-codigo'); return c && c.textContent.includes('linha 2') && c; });
          return { mesmaMidia, novoCodigo: cod !== window.__cod };
        """)
        self.assertTrue(r["mesmaMidia"])
        self.assertTrue(r["novoCodigo"], "o arquivo que mudou no disco foi lido de novo")

    def test_trocar_de_conversa_so_confere_a_marca(self):
        buf = io.BytesIO()
        with wave.open(buf, "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(8000)
            w.writeframes(b"\0\0" * 8000)
        p = self.escrever("som.wav", buf.getvalue())
        conv = banco.criar_conversa("code", pasta=self.raiz)
        _checkpoint(conv["id"], p, b"antes")
        # a primeira mensagem de um trabalho novo cria a conversa: o contexto passa de sem conversa para ela
        r = self.na_tela("const CONV = " + json.dumps(conv["id"]) + ";\n" + r"""
          A.abrirArquivo('som.wav');
          const midia = await esperar(() => vis()?.querySelector('.vis-midia audio'));
          await esperar(() => midia.readyState >= 1);
          const botao = () => vis().querySelector('.vis-mudancas');
          const antes = !!botao();
          P.definirContexto({ modo: 'code', pasta: RAIZ, conversaId: CONV });
          await esperar(botao);
          const mesmaComMarca = vis().querySelector('.vis-midia audio') === midia;
          P.definirContexto({ modo: 'code', pasta: RAIZ, conversaId: null });
          await esperar(() => !botao());
          return { antes, mesmaComMarca, mesmaSemMarca: vis().querySelector('.vis-midia audio') === midia };
        """)
        self.assertFalse(r["antes"])
        self.assertTrue(r["mesmaComMarca"], "a marca do Claude apareceu sem recarregar o áudio")
        self.assertTrue(r["mesmaSemMarca"])

    def test_clique_duplo_na_pasta_abre_como_no_explorador(self):
        self.escrever("pasta/dentro.txt", "x")
        r = self.na_tela(r"""
          P.abrirPainel('arquivos');
          const linha = () => arq()?.querySelector('.arq-linha[data-caminho="pasta"]');
          await esperar(linha);
          const evento = (tipo, detail) => linha().dispatchEvent(new MouseEvent(tipo, { bubbles: true, detail }));
          evento('click', 1);
          await dormir(150);
          evento('click', 2);
          evento('dblclick', 2);
          await esperar(() => arq().querySelector('.arq-linha[data-caminho="pasta/dentro.txt"]'));
          await dormir(300);
          const duplo = linha().getAttribute('aria-expanded');
          evento('click', 1);
          await dormir(300);
          return { duplo, simples: linha().getAttribute('aria-expanded') };
        """)
        self.assertEqual(r["duplo"], "true", "o clique duplo deixa a pasta aberta")
        self.assertEqual(r["simples"], "false", "um clique continua abrindo e fechando")

    def test_contagem_da_busca_no_formato_do_idioma(self):
        self.escrever("muitos.txt", "a\n" * 20005)
        r = self.na_tela(r"""
          const I = await import('/js/i18n.js');
          await I.carregarIdioma('en');
          try {
            A.abrirArquivo('muitos.txt');
            await esperar(() => vis()?.querySelector('.vis-codigo'));
            vis().querySelector('.vis-corpo').dispatchEvent(new KeyboardEvent('keydown', { key: 'f', ctrlKey: true, bubbles: true }));
            const campo = await esperar(() => { const c = vis().querySelector('.vis-busca-campo'); return c && !c.closest('.vis-busca').hidden && c; });
            campo.value = 'a';
            campo.dispatchEvent(new Event('input', { bubbles: true }));
            return await esperar(() => { const x = vis().querySelector('.vis-busca-contagem').textContent; return x.includes('+') && x; }, 30000);
          } finally { await I.carregarIdioma('pt-BR'); }
        """)
        self.assertEqual(r, "1 of 20,000+")


class TestAcentosEOrdem(unittest.TestCase):
    def test_ordem_natural(self):
        nomes = ["arquivo10", "Arquivo2", "arquivo1", "beta", "Álbum"]
        self.assertEqual(sorted(nomes, key=explorador._chave_natural)[:3], ["arquivo1", "Arquivo2", "arquivo10"])

    def test_decodificar_direto(self):
        self.assertIsNone(explorador.decodificar(b"\x00\x01\x02abc"))
        self.assertEqual(explorador.decodificar("olá".encode("utf-8") + b"\xc3", cortado=True), ("olá", "utf-8"))
        self.assertEqual(explorador.decodificar(b""), ("", "utf-8"))


if __name__ == "__main__":
    unittest.main()
